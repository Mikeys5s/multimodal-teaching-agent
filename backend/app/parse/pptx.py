"""PPTX 解析（归属：P1）。SPEC §5.1。

一张幻灯片 = 一"页"
-------------------
PPTX 里没有"页"，但幻灯片正是它的天然单位。所以 `page_no` 取 **1-based 幻灯片
序号**（第 1 张 = 1）、`page_count` = 幻灯片张数 —— 与 PDF 的页码同口径：
用户说"第 3 页"、`slides[2]` 与库里的 `page_no=3` 三者直接对得上，中间不需要
任何换算。这对 A1-6（块可定位到页码）是免费的：动画一停，观众看到的就是那一页。

`line_start` / `line_end` 为什么是 NULL
---------------------------------------
那两个字段的语义是"**页内**第几行"，服务对象是 PDF —— PDF 的文本层按行给出
坐标，`page_no + line_start/line_end` 能框出一小段原文。PPTX 里没有"行"这个
对象：文本框里的换行是**排版结果**（字号、框宽一变就重排），随手 resize 一下
行号就全变了。给一个会自己漂移的行号会让溯源指向不存在的行，比没有行号更糟，
与 DOCX 同一取舍（见 `docx.py` 模块 docstring）。所以一律 NULL，溯源靠
「幻灯片号 + 全局块序号」（`markdown.block_anchor`）。

备注（notes_slide）为什么是 `other` 块 + `> 备注：` 前缀
--------------------------------------------------------
备注是讲课时真正要说的话，对"讲什么、怎么讲"价值很高，扔掉可惜；但它**不是**
幻灯片上学生能看到的内容，混进正文会污染知识点抽取。两头都要，就只能显式区分，
于是两层一起做：

  · `block_type="other"` —— 机器一眼能滤掉（8 个枚举值里 `other` 就是"不属于
    其余七类"的语义，备注正合适）；
  · `content_md` 前缀 `> 备注：` —— 人和 LLM 在 Markdown 里也一眼看得出来，
    而且 `>` 是标准 Markdown 引用语法，渲染出来就是引用块。

只做其一都会留个坑：只改 `block_type` 的话，备注拼进 Markdown 后与正文无从分辨；
只加前缀的话，下游按 `block_type` 过滤时会把备注当正文。

标题层级怎么判
--------------
PPTX 的标题占位符**没有** Word 那种大纲级别（`outlineLvl`），所以默认记 **1 级**
—— 幻灯片标题在结构上就是一级。唯一的例外是段落自带的缩进级别
（`<a:pPr lvl>`，0-based，0 = 顶层）：有就用 `min(lvl + 1, 6)`，越界夹到 6
（DB 有 1–6 的 CHECK）。绝大多数文件里 `lvl` 就是 0，于是结果就是 1 级。

正文顺序按 XML 声明顺序走
-------------------------
`slide.shapes.title` / `slide.placeholders` 是**派生的便利访问器**：它们给出的
顺序不是版面声明顺序，而且会整体漏掉非占位符的文本框（用空白版式 + 文本框搭的
PPT 就是这种）。这里统一遍历 `slide.shapes`，并**递归展开组合形状**（group）——
组里的文本框和表格在 PPT 里很常见，不递归就是静默丢内容（解析显示成功、材料
却残缺且无人知道），属于最坏的一类失败。

已知不做的事
------------
图片 / 图表 / SmartArt 里的文字不抽取（本批次没有多模态能力，与"扫描版 PDF
只识别不解析"同一取舍）；不做页眉页脚去重（PPT 没有这个概念）。
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from pathlib import Path

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.presentation import Presentation as PptxPresentation
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.table import Table
from pptx.text.text import TextFrame

from app.core.errors import ApiError, ErrorCode
from app.parse.blocks import (
    HEADING_LEVEL_MAX,
    HEADING_LEVEL_MIN,
    ParsedBlock,
    ParsedDocument,
    UncertainNote,
)

#: OLE2 复合文档头。加密的 .pptx 与旧版 .ppt 都是这个头，两者对用户的处置动作
#: 都是"用 PowerPoint 另存为未加密的 .pptx"，所以合并成一条提示（与 docx.py 同）。
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

#: 备注块的内容前缀（见模块 docstring）。
_NOTES_PREFIX = "> 备注："

#: 认作"标题"的占位符类型。竖排标题只是排版方向不同，同样是标题。
_TITLE_PLACEHOLDER_TYPES = (
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
)


# ---------------------------------------------------------------------------
# 文本
# ---------------------------------------------------------------------------


def _clean_text(text: str) -> str:
    """按行归一：行长首尾去空白、行内空白压成单空格、丢掉空行。

    保留换行是**有意**的：PPT 里"一行一行"排出来的清单项、公式的两边，对下游
    有价值，而 Markdown 段落里的单换行渲染出来仍是同一段，不会破坏结构。
    `splitlines()` 顺带吃掉 `\\v`（PPT 内部软换行的编码）与 `\\r`。
    """
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _one_line(text: str) -> str:
    """压成一行。

    标题与备注必须单行：标题要交给 `render_block` 拼 `# ` 前缀，内容里带换行
    会把一个标题渲染成"标题 + 正文"两个块；备注已经用 `>` 前缀标了来源，
    再折行只是噪声。
    """
    return " ".join(text.split())


# ---------------------------------------------------------------------------
# 形状遍历
# ---------------------------------------------------------------------------


def _iter_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """按声明顺序产出形状；组合形状递归展开（组里的文字不能丢）。"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _is_title(shape: BaseShape) -> bool:
    """是不是标题占位符。非占位符的文本框一律当正文 —— 不猜。"""
    if not shape.is_placeholder:
        return False
    return shape.placeholder_format.type in _TITLE_PLACEHOLDER_TYPES


def _title_level(text_frame: TextFrame) -> int:
    """标题层级：取最靠外那段的大纲缩进级别 + 1，夹到 1–6（见模块 docstring）。"""
    levels = [p.level for p in text_frame.paragraphs if p.text.strip()]
    if not levels:
        return HEADING_LEVEL_MIN
    return min(min(levels) + 1, HEADING_LEVEL_MAX)


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------


def _clean_cell(text: str) -> str:
    """单元格文本 → 一行 Markdown 单元格。

    与 `docx.py::_clean_cell` 同一口径：压掉所有空白、转义 `|`（不转义的话
    单元格里的 `|` 会把表格撕成多列，整张表就塌了）。
    """
    return " ".join(text.split()).replace("|", "\\|")


def _table_to_md(table: Table) -> str:
    """PPT 表格 → Markdown 表格；整表无内容时返回空串（调用方据此不产块）。

    与 `docx.py::_table_to_md` 同一口径：
      · 首行当表头（Markdown 表格语法要求有表头行，PPT 里的表格首行也绝大多数
        就是表头）；
      · 各行列数不一致时右侧补空单元格 —— 列数不齐的 Markdown 表格在多数
        渲染器里会直接塌掉；
      · **整表无内容 → 空串**。否则 `|  |` + 分隔行 `strip()` 之后仍含 `|`，
        会躲过 `ParsedDocument` 的空白块过滤，变成一个纯噪声块。

    合并单元格：python-pptx 对被合并"吃掉"的格子给出 `is_spanned=True` 且
    `text` 为空 —— 这里只写合并起点那一格，续格留空，内容不会在 Markdown 里
    被复制成 n 份（Markdown 没有 colspan / rowspan，这是唯一不失真的表达）。

    与 DOCX 的差别：PPT 的表格单元格里**只能有文本**，不存在嵌套表格，
    所以不需要 docx 那套递归展平。
    """
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append(["" if cell.is_spanned else _clean_cell(cell.text) for cell in row.cells])

    if not rows:
        return ""
    width = max(len(r) for r in rows)
    if width == 0:
        return ""
    padded = [r + [""] * (width - len(r)) for r in rows]
    if not any(cell for row in padded for cell in row):
        return ""
    header = "| " + " | ".join(padded[0]) + " |"
    sep = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(r) + " |" for r in padded[1:]]
    return "\n".join([header, sep, *body])


# ---------------------------------------------------------------------------
# 幻灯片 → 块
# ---------------------------------------------------------------------------


def _slide_blocks(slide: Slide, page_no: int) -> list[ParsedBlock]:
    """一张幻灯片上的形状 → 块，顺序即声明顺序（见模块 docstring）。

    一个正文段落产出一块（与 DOCX 逐段产块的粒度一致）：项目符号列表里每条
    都是一个独立的知识单元，合成一大块会让溯源只能指向"整页"。
    """
    blocks: list[ParsedBlock] = []
    for shape in _iter_shapes(slide.shapes):
        if shape.has_table:
            md = _table_to_md(shape.table)
            if md:
                blocks.append(ParsedBlock(block_type="table", content_md=md, page_no=page_no))
            continue
        if not shape.has_text_frame:
            continue  # 图片 / 图表 / 连接线：本批次不抽（见模块 docstring）
        if _is_title(shape):
            text = _one_line(shape.text_frame.text)
            if text:
                blocks.append(
                    ParsedBlock(
                        block_type="heading",
                        content_md=text,
                        page_no=page_no,
                        heading_level=_title_level(shape.text_frame),
                    )
                )
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = _clean_text(paragraph.text)
            if text:
                blocks.append(ParsedBlock(block_type="paragraph", content_md=text, page_no=page_no))
    return blocks


def _notes_blocks(slide: Slide, page_no: int) -> list[ParsedBlock]:
    """演讲者备注 → `other` 块（带 `> 备注：` 前缀，见模块 docstring）。

    先用 `has_notes_slide` 探一下：`slide.notes_slide` 会在**没有备注页时凭空
    创建一张**，那是写操作 —— 解析器只读，不该改动文档。
    """
    if not slide.has_notes_slide:
        return []
    text_frame = slide.notes_slide.notes_text_frame
    if text_frame is None:  # 备注母版里没有正文占位符（少见，如实跳过）
        return []
    blocks: list[ParsedBlock] = []
    for paragraph in text_frame.paragraphs:
        text = _one_line(paragraph.text)
        if text:
            blocks.append(
                ParsedBlock(
                    block_type="other",
                    content_md=f"{_NOTES_PREFIX}{text}",
                    page_no=page_no,
                )
            )
    return blocks


# ---------------------------------------------------------------------------
# 打开与解析
# ---------------------------------------------------------------------------


def _open_pptx(path: str | Path) -> PptxPresentation:
    """打开 PPTX；空文件 / 加密 / 旧版 .ppt / 损坏全部转成中文 `ApiError`。"""
    # 0 字节先单独判：python-pptx 会把空文件报成"损坏"，措辞与 PDF/DOCX 的
    # "空文件"分支不一致 —— 同一类输入（上传时被截断成 0 字节）在几个格式下
    # 应该给同一句人话，用户才知道该做什么。
    try:
        is_empty = Path(path).stat().st_size == 0
    except OSError:
        is_empty = False  # 读不到就交给下面的打开流程，按损坏处理
    if is_empty:
        raise ApiError(
            ErrorCode.INVALID_PARAM,
            "这个 PPT 文件是空文件（0 字节），里面没有内容，请重新导出 .pptx 后再上传",
        )

    try:
        return pptx.Presentation(str(path))
    except Exception as exc:  # noqa: BLE001 —— python-pptx 的异常类型不稳定，见下
        # 与 docx.py 同一口径：不猜异常类型（python-pptx 会把 zipfile.BadZipFile /
        # KeyError / ValueError 直接往外抛，版本间还不一致），改按**文件头**判断：
        #   · OLE2 头 → 加密的 pptx 或旧版 .ppt；
        #   · 其余非 zip → 损坏。
        head = b""
        try:
            with open(path, "rb") as fh:
                head = fh.read(8)
        except OSError:
            pass  # 读不到就按损坏处理，下面统一报错

        if head.startswith(_OLE2_MAGIC):
            raise ApiError(
                ErrorCode.INVALID_PARAM,
                "这个 PPT 文件已加密，或是旧版 .ppt 格式，无法解析，"
                "请用 PowerPoint 另存为未加密的 .pptx 再上传",
            ) from exc
        raise ApiError(
            ErrorCode.INVALID_PARAM,
            "这个 PPT 文件已损坏或不是有效的 .pptx，无法打开，请重新导出后再上传",
        ) from exc


def parse_pptx(path: str | Path) -> ParsedDocument:
    """解析 PPTX：一张幻灯片 = 一页，产出按放映顺序排列的块。

    · 标题占位符 → `heading` + `heading_level`（默认 1，见模块 docstring）
    · 正文文本框 / 项目符号 → `paragraph`（一个段落一块）
    · 表格 → `table`（Markdown 表格；合并格内容只写一次、空表不产块）
    · 演讲者备注 → `other`，`content_md` 带 `> 备注：` 前缀
    · `page_no` = 1-based 幻灯片序号（**每块都有**），`page_count` = 幻灯片张数
    · `line_start` / `line_end` 一律 NULL（PPTX 没有行号概念，见模块 docstring）
    · 空白幻灯片 / 空文本框 / 空表格**不产块**（`ParsedDocument` 还会再过一遍
      空白块，那是唯一过滤点）
    · 一张幻灯片都没有的空 PPT 与 0 字节文件一样按用户错误拒绝：它们都会让材料
      显示"解析成功"却零产物，是最不该静默的那种失败。
    · 失败一律抛中文 `ApiError`（加密、旧版 .ppt、损坏、空文件、无幻灯片）。
    """
    presentation = _open_pptx(path)

    if len(presentation.slides) == 0:
        raise ApiError(
            ErrorCode.INVALID_PARAM,
            "这个 PPT 里没有任何幻灯片，没有可解析的内容，请确认后重新导出 .pptx 再上传",
        )

    blocks: list[ParsedBlock] = []
    for index, slide in enumerate(presentation.slides):
        page_no = index + 1  # 1-based：幻灯片序号就是页码
        blocks.extend(_slide_blocks(slide, page_no))
        # 备注紧跟在它所属的那一页后面 —— 与"这页讲什么"挨着，人工核对时顺眼。
        blocks.extend(_notes_blocks(slide, page_no))

    notes: list[UncertainNote] = []
    if not any(b.is_heading for b in blocks):
        # 与 docx.py 同一口径：骨架推不出来就老实说，不硬造节。
        notes.append(
            UncertainNote(
                kind="ambiguous_structure",
                severity="medium",
                message=(
                    "这份 PPT 里没有检测到任何标题占位符，已把整份材料作为一节处理，"
                    "建议人工确认章节结构后再抽取知识点。"
                ),
            )
        )

    return ParsedDocument(
        source_type="pptx",
        blocks=blocks,
        page_count=len(presentation.slides),
        parse_method="text_extract",
        uncertain_notes=notes,
    )
