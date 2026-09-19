"""DOCX 解析（归属：P1）。SPEC §5.1。

`page_no` 为什么是 NULL
-----------------------
DOCX **没有页的概念** —— 分页是 Word 打开时按纸张大小、字号、打印机驱动
实时算出来的，同一个文件在两台机器上页数都可能不同。所以这里老老实实写
NULL（data-model 允许），不按段落数"估"一个页码：估出来的页码会让溯源指到
不存在的位置，比没有页码更糟。

这正是 docs/extraction-channel.md §6 问题 2 的结论 —— 溯源以原文片段为准，
页码只是增强。网页版教材连页码都没有，靠「章节号 + 块序号」照样能定位。

正文顺序：必须走 body 的 XML 子节点
-----------------------------------
`document.paragraphs` 与 `document.tables` 是**两个独立的列表**，各自内部有序、
但彼此之间的先后关系丢了。教材里"表格夹在段落中间"是常态，用那两个列表拼
出来的 Markdown 会张冠李戴。所以这里直接遍历 `body` 的子元素。
单元格内部同理（见 `_cell_content`）：直接段落与嵌套表也要按 XML 顺序取。

`page_no` / `line_start` / `line_end` 三个字段一起为 NULL
--------------------------------------------------------
不只是 `page_no`：DOCX **没有页、也没有"页内行号"这两个概念**。
   · `page_no`：分页是 Word 打开时按纸张/字号/打印机驱动实时算的，同一文件在
     两台机器上页数都可能不同；
   · `line_start` / `line_end`：那两个字段的语义是"**页内**第几行"（PDF 用它们
     与 `page_no` 一起定位），没有页就无所谓"页内"，单独给一个全文行号会让
     溯源指向一个不存在的坐标系。
所以三者一律 NULL（data-model 允许）。溯源在这一侧靠「章节号 + 块序号」
（`markdown.block_anchor`）落地 —— 网页版教材连页码都没有，靠它照样能定位。
刻意**不**按段落数"估"页码/行号：估出来的位置会让溯源指到不存在的地方，
比没有位置更糟（docs/extraction-channel.md §6 问题 2）。
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.core.errors import ApiError, ErrorCode
from app.parse.blocks import ParsedBlock, ParsedDocument, UncertainNote, split_heading_number

#: Word 内建标题样式的名字。中文版 Word 里可能显示成「标题 1」，
#: python-docx 取到的是 style.name，两种都要认。
_HEADING_STYLE_RE = re.compile(r"^(?:Heading|标题)\s*([1-9])$", re.IGNORECASE)

#: 编号法判标题的长度上限。比 PDF 宽松一点：Word 的段落本身就是一段，
#: 不存在"多行块"的歧义。
_HEADING_MAX_CHARS = 100

#: OLE2 复合文档头。加密的 .docx 与旧版 .doc 都是这个头 —— 两者对用户的
#: 处置动作都是"另存为未加密的 .docx"，所以合并成一条提示。
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _heading_level_of(paragraph: Paragraph) -> int | None:
    """段落是不是标题；是的话给出 1–6 级。

    三个来源，按可信度排序：
      1. 样式名（`Heading 2` / `标题 2`）—— 作者显式标了，最可信；
      2. `w:pPr/w:outlineLvl`（大纲级别 0–8 → 1–9 级）—— 有些模板用大纲级别
         而不是标题样式；
      3. 编号法兜底，但**只认多级编号**（`3.1`、`2.3.4`）。
    第 3 条刻意不收单级编号（`1.`）：在 Word 里那几乎总是编号列表项，
    收进来会把列表项变成"章"，整份骨架就歪了。
    """
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    m = _HEADING_STYLE_RE.match(name.strip())
    if m:
        return min(int(m.group(1)), 6)

    ppr = paragraph._p.pPr
    if ppr is not None and ppr.outlineLvl is not None:
        raw = ppr.outlineLvl.val
        if isinstance(raw, int) and raw < 9:  # 9 = 正文
            return min(raw + 1, 6)

    text = paragraph.text.strip()
    if len(text) <= _HEADING_MAX_CHARS:
        number = split_heading_number(text)
        if number is not None and number.depth >= 2:
            return min(number.depth, 6)
    return None


def _clean_cell(text: str) -> str:
    """单元格文本 → 一行 Markdown 单元格内容。"""
    return " ".join(text.split()).replace("|", "\\|")


def _row_cells_with_gaps(row: Any) -> Iterator[_Cell | None]:
    """按真实列数逐格产出单元格；**横向合并**的续格产出 `None`。

    python-docx 的 `row.cells` 对跨 n 列的合并单元格会把**同一个 `_Cell`
    对象返回 n 次**（列数因此保持为表格的真实列数，这是好事），于是
    「跨三列」这句在 Markdown 里会被复制 n 份。Markdown 没有 colspan，
    正确处理是"内容写第一格、其余留空" —— 所以这里只在第一格给出单元格，
    续格给出 `None`，由调用方决定"留空"（Markdown）还是"跳过"（展平文本）。
    """
    seen: set[int] = set()
    for cell in row.cells:
        if id(cell) in seen:
            yield None
            continue
        seen.add(id(cell))
        yield cell


def _iter_children(parent_xml: Any, parent: Any) -> Iterator[tuple[str, Any]]:
    """按 XML 顺序遍历一个容器的子元素：段落与表格交错出现，顺序必须保持原样。

    `body` 与**单元格**都用它：两处的 `paragraphs` / `tables` 都是彼此独立的
    列表，"表格夹在段落中间"的排版只有按 XML 子节点走才不会张冠李戴。
    """
    for child in parent_xml.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield "tbl", Table(child, parent)


def _flatten_table(table: Table) -> tuple[str, int]:
    """把一张表（含更深层的嵌套表）压成**一行纯文本**；返回 `(文本, 嵌套表数)`。

    格式：单元格之间用空格、行之间用 ` / `。
    刻意**不用** Markdown 的 `|`：这段文本要放进外层表格的单元格里，带 `|`
    会把外层表格撕成多列。展平是"不丢内容"与"不破坏外层表格列数"之间唯一
    可行的折中 —— 嵌套层次本身确实丢了，所以由调用方记一条中文存疑说明。
    """
    rows: list[str] = []
    nested_tables = 0
    for row in table.rows:
        cells: list[str] = []
        for cell in _row_cells_with_gaps(row):
            if cell is None:  # 合并的续格：展平文本里不再重复一遍内容
                continue
            text, cell_nested = _cell_content(cell)
            nested_tables += cell_nested
            if text:
                cells.append(text)
        if cells:
            rows.append(" ".join(cells))
    return " / ".join(rows), nested_tables


def _cell_content(cell: _Cell) -> tuple[str, int]:
    """单元格的**全部**文本 + 它里面的嵌套表数量。

    为什么不能直接用 `cell.text`：python-docx 的 `Cell.text` 只拼接该单元格的
    **直接段落**（`cell.paragraphs`），内嵌表格（`cell.tables`）被整段跳过 ——
    "外层单元格里塞了一张表"这种排版会**静默丢内容**：输出的 Markdown 看起来
    是一张完整的表，内层内容却一个字都没有，存疑说明也是空的。这是最坏的一类
    失败（用户以为解析成功，实际材料残缺且无人知道），所以这里递归展开：
    直接段落与嵌套表按 XML 顺序取，嵌套表展平后接在同一格文本里。
    """
    parts: list[str] = []
    nested_tables = 0
    for kind, part in _iter_children(cell._tc, cell):
        if kind == "p":
            text = part.text.strip()
            if text:
                parts.append(text)
        else:
            nested_tables += 1
            text, inner_nested = _flatten_table(part)
            nested_tables += inner_nested  # 内层还可以再嵌，一起上报
            if text:
                parts.append(text)
    return " ".join(parts), nested_tables


def _table_to_md(table: Table) -> tuple[str, int]:
    """把 Word 表格渲染成 Markdown 表格；返回 `(Markdown, 嵌套表数量)`。

    第一行当表头（Markdown 表格语法要求有表头行）。教材里的表格绝大多数
    首行就是表头；遇到不是的，Markdown 本身表达不了"无表头表格"，只能这样。
    合并单元格会让某些行的单元格数少于最长行，用空单元格补齐，保证列数一致 ——
    列数不齐的 Markdown 表格在多数渲染器里会直接塌掉。

    三个必须在**这里**处理掉的坑：
      · 嵌套表：内容由 `_cell_content` 递归展平写进单元格，绝不静默丢弃；
      · 横向合并：同一格只写一次，其余留空（见 `_row_cells_with_gaps`）；
      · 整表无内容：返回空串，让调用方不产块 —— `|  |` + 分隔行 `strip()` 之后
        仍含 `|`，会躲过 `ParsedDocument` 的空白块过滤，变成一个纯噪声块。
    """
    rows: list[list[str]] = []
    nested_tables = 0
    for row in table.rows:
        cells: list[str] = []
        for cell in _row_cells_with_gaps(row):
            if cell is None:
                cells.append("")  # 横向合并的续格：留空，列数不变
                continue
            text, cell_nested = _cell_content(cell)
            nested_tables += cell_nested
            cells.append(_clean_cell(text))
        rows.append(cells)

    if not rows:
        return "", nested_tables
    width = max(len(r) for r in rows)
    if width == 0:
        return "", nested_tables
    padded = [r + [""] * (width - len(r)) for r in rows]
    if not any(cell for row in padded for cell in row):
        return "", nested_tables
    header = "| " + " | ".join(padded[0]) + " |"
    sep = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(r) + " |" for r in padded[1:]]
    return "\n".join([header, sep, *body]), nested_tables


def _iter_body_parts(document: DocxDocument) -> Iterator[tuple[str, Any]]:
    """按 XML 顺序遍历 body：段落与表格交错出现，顺序必须保持原样。"""
    yield from _iter_children(document.element.body, document)


def _open_docx(path: str | Path) -> DocxDocument:
    """打开 DOCX；空文件 / 加密 / 损坏 / 旧版 .doc 全部转成中文 `ApiError`。"""
    # 0 字节先单独判：python-docx 会把空文件抛成 BadZipFile，于是报"已损坏"，
    # 与 PDF 分支的"空文件"措辞不一致 —— 同一类输入（上传时被截断成 0 字节）
    # 在两个格式下应该给同一句人话，用户才知道该做什么。
    try:
        is_empty = Path(path).stat().st_size == 0
    except OSError:
        is_empty = False  # 读不到就交给下面的打开流程，按损坏处理
    if is_empty:
        raise ApiError(
            ErrorCode.INVALID_PARAM,
            "这个 Word 文件是空文件（0 字节），里面没有内容，请重新导出 .docx 后再上传",
        )

    try:
        return docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 —— python-docx 的异常类型不稳定，见下
        # python-docx 把 zipfile.BadZipFile / KeyError / ValueError 直接往外抛，
        # 版本间还不一致，所以这里按**文件头**判断，比猜异常类型靠谱：
        #   · OLE2 头 → 加密的 docx 或旧版 .doc
        #   · 其余非 zip → 损坏
        head = b""
        try:
            with open(path, "rb") as fh:
                head = fh.read(8)
        except OSError:
            pass  # 读不到就按损坏处理，下面统一报错

        if head.startswith(_OLE2_MAGIC):
            raise ApiError(
                ErrorCode.INVALID_PARAM,
                "这个 Word 文件已加密，或是旧版 .doc 格式，无法解析，"
                "请先用 Word 另存为未加密的 .docx 再上传",
            ) from exc
        raise ApiError(
            ErrorCode.INVALID_PARAM,
            "这个 Word 文件已损坏或不是有效的 .docx，无法打开，请重新导出后再上传",
        ) from exc


def parse_docx(path: str | Path) -> ParsedDocument:
    """解析 DOCX，产出按文档顺序排列的块。

    · 标题 → `heading` + `heading_level`（1–6）
    · 表格 → `table`（Markdown 表格；嵌套表展平，见 `_table_to_md`）
    · 其余非空段落 → `paragraph`
    · `page_no` / `line_start` / `line_end` 一律 NULL（DOCX 无页、也无页内行号，
      见模块 docstring）
    · 空白段落在这里就被丢掉，`ParsedDocument` 还会再过一遍（唯一过滤点）
    """
    document = _open_docx(path)

    blocks: list[ParsedBlock] = []
    nested_tables = 0
    for kind, part in _iter_body_parts(document):
        if kind == "p":
            if not isinstance(part, Paragraph):  # pragma: no cover - 防御性分支
                continue
            text = part.text.strip()
            if not text:
                continue
            level = _heading_level_of(part)
            if level is None:
                blocks.append(ParsedBlock(block_type="paragraph", content_md=text))
            else:
                blocks.append(
                    ParsedBlock(block_type="heading", content_md=text, heading_level=level)
                )
        else:
            if not isinstance(part, Table):  # pragma: no cover - 防御性分支
                continue
            md, nested = _table_to_md(part)
            nested_tables += nested
            if md.strip():
                blocks.append(ParsedBlock(block_type="table", content_md=md))

    notes: list[UncertainNote] = []
    if not any(b.is_heading for b in blocks):
        notes.append(
            UncertainNote(
                kind="ambiguous_structure",
                severity="medium",
                message=(
                    "这份 Word 文档里没有检测到任何标题样式，已把整份材料作为一节处理，"
                    "建议人工确认章节结构后再抽取知识点。"
                ),
            )
        )
    if nested_tables:
        # 内容是**保住了**的（已展平写进外层单元格），丢的是表格层次 ——
        # 所以严重度 low，但不能不吭声：下游按行列取值时会读到展平后的串。
        notes.append(
            UncertainNote(
                kind="ambiguous_structure",
                severity="low",
                message=(
                    f"这份 Word 文档里有 {nested_tables} 个嵌套表格（表格里还套着表格）。"
                    "Markdown 表达不了嵌套结构，已把内层表格的内容展平后写进外层单元格"
                    "（单元格内用空格分隔、行之间用「/」分隔），表格层次需要人工核对。"
                ),
            )
        )

    return ParsedDocument(
        source_type="docx",
        blocks=blocks,
        page_count=None,  # DOCX 无页概念
        parse_method="text_extract",
        uncertain_notes=notes,
    )
