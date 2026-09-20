"""PPTX 解析的测试（归属：P1）。SPEC §5.1。

守的是这几条契约：
  · 标题占位符 → `block_type="heading"` 且 `heading_level` 合法（1–6）
  · 正文文本框 / 项目符号 → `paragraph`
  · 表格 → `table`（Markdown 表格；横向合并格内容只写一次、整表无内容不产块）
  · 演讲者备注 → 产出块，且**能与正文区分**（`other` + `> 备注：` 前缀）
  · `page_no` = 1-based **幻灯片序列**，每块都有（与 DOCX 的"一律 NULL"相反）；
    空白幻灯片照样占一个页码
  · `line_start` / `line_end` 一律 None —— PPTX 里没有"页内行号"这个坐标系
  · `seq` 是整份材料内的全局序号，从 0 开始连续
  · 块锚点里的 block_id 与 `app.models.ids.block_id()` 现算结果一致
  · 坏输入（空文件 / 无幻灯片 / 损坏 / 非 pptx 字节）→ 中文 `ApiError`
  · A2-7 同输入可复现

夹具全部用 python-pptx **现场生成**，不依赖仓库外的素材文件 —— 素材文件进不了
版本库，靠它们写的测试在队友机器上必然红（与 `test_parse_docx.py` 同一条纪律）。
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.presentation import Presentation as PptxPresentation
from pptx.util import Inches
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ApiError, ErrorCode
from app.models._common import BLOCK_TYPES
from app.models.ids import block_id, hash_bytes, material_id
from app.models.material import Material, MaterialBlock
from app.parse import parse_material, parse_pptx, to_markdown
from app.parse.persist import persist_blocks

_HAN_RE = re.compile(r"[\u4e00-\u9fff]")

#: 块锚点行：`<!-- block: blk_a1b2c3d4_00007 -->`
_ANCHOR_RE = re.compile(r"<!-- block: (blk_\w+) -->")

#: 与 test_ids.py 同口径：material_id 只取哈希前 8 位。
MAT_ID = material_id(hash_bytes(b"pptx fixture"))


# ---------------------------------------------------------------------------
# 夹具生成
# ---------------------------------------------------------------------------


def _layout(prs: PptxPresentation, name: str):
    """按版式名取版式 —— 比记版面索引稳（不同模板的索引顺序不一样）。"""
    return next(lay for lay in prs.slide_layouts if lay.name == name)


def _blank_slide(prs: PptxPresentation):
    """空白版式：没有任何占位符，只能靠文本框 / 表格，正好覆盖"非占位符形状"。"""
    return prs.slides.add_slide(_layout(prs, "Blank"))


def make_deck(path: Path) -> Path:
    """4 张幻灯片，覆盖标题 / 项目符号 / 文本框 / 合并单元格表格 / 空白 / 空表。"""
    prs = Presentation()

    # 第 1 张：标题占位符 + 正文占位符（含二级项目符号）+ 演讲者备注
    slide1 = prs.slides.add_slide(_layout(prs, "Title and Content"))
    slide1.shapes.title.text = "第1章 网络层"
    body = slide1.placeholders[1].text_frame
    body.text = "路由是把分组从源送到目的的过程。"
    body.add_paragraph().text = "路由器根据路由表决定下一跳。"
    nested = body.add_paragraph()
    nested.text = "最长前缀匹配"
    nested.level = 1
    notes = slide1.notes_slide.notes_text_frame
    notes.text = "备注第一段：先讲清楚“转发”与“路由”的区别。"
    notes.add_paragraph().text = "备注第二段：留 2 分钟做课堂提问。"

    # 第 2 张：空白版式 —— 文本框 + 带横向合并的表格（都不是占位符）
    slide2 = _blank_slide(prs)
    box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    box.text_frame.text = "补充说明：这张表对比两种协议。"
    table = slide2.shapes.add_table(3, 3, Inches(0.5), Inches(2), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "协议"
    table.cell(0, 1).text = "作用"
    table.cell(0, 2).text = "层次"
    table.cell(1, 0).text = "IP"
    table.cell(1, 1).merge(table.cell(1, 2))  # 合并后内容写在起点格 (1,1)
    table.cell(1, 1).text = "跨两列"
    table.cell(2, 0).text = "TCP"
    table.cell(2, 1).text = "可靠传输"
    table.cell(2, 2).text = "传输层"

    # 第 3 张：完全空白 —— 不产块，但页码要占住（第 4 张仍是第 4 页）
    _blank_slide(prs)

    # 第 4 张：只有一张全空的表格 —— 同样不产块（不许留下 `|  |` 这种噪声块）
    slide4 = _blank_slide(prs)
    slide4.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(2), Inches(1))

    prs.save(str(path))
    return path


@pytest.fixture
def fixture_pptx(tmp_path: Path) -> Path:
    return make_deck(tmp_path / "讲义.pptx")


def _assert_chinese_error(err: ApiError, keyword: str) -> None:
    """报错文案必须是"能直接给用户看的中文"，并给出下一步动作。"""
    assert err.code == ErrorCode.INVALID_PARAM
    assert err.status_code == 400
    assert _HAN_RE.search(err.message), f"报错文案里没有中文：{err.message!r}"
    assert keyword in err.message, f"文案没说到 {keyword!r}：{err.message!r}"
    assert "请" in err.message, f"报错文案没有给出下一步动作：{err.message!r}"


# ---------------------------------------------------------------------------
# 文档级契约
# ---------------------------------------------------------------------------


def test_pptx_source_type_and_slide_count(fixture_pptx: Path) -> None:
    """PPTX 走文本抽取；`page_count` = 幻灯片张数（空幻灯片也算一张）。"""
    doc = parse_pptx(fixture_pptx)
    assert doc.source_type == "pptx"
    assert doc.parse_method == "text_extract"
    assert doc.page_count == 4
    assert doc.blocks


def test_every_block_type_is_in_the_enum(fixture_pptx: Path) -> None:
    doc = parse_pptx(fixture_pptx)
    types = {b.block_type for b in doc.blocks}
    assert types <= set(BLOCK_TYPES)
    assert types == {"heading", "paragraph", "table", "other"}


# ---------------------------------------------------------------------------
# 标题 / 层级
# ---------------------------------------------------------------------------


def test_title_placeholder_becomes_a_level_1_heading(fixture_pptx: Path) -> None:
    """★ 标题占位符 → `heading`；PPTX 没有大纲级别，默认 1 级。"""
    doc = parse_pptx(fixture_pptx)
    headings = doc.headings
    assert [h.content_md for h in headings] == ["第1章 网络层"]
    assert headings[0].heading_level == 1
    assert headings[0].page_no == 1


def test_heading_level_is_legal_and_none_for_non_heading_blocks(fixture_pptx: Path) -> None:
    """非 heading 块带上 level 会让下游凭空多出一节，必须在解析层就为 None。"""
    doc = parse_pptx(fixture_pptx)
    assert doc.headings
    for block in doc.blocks:
        if block.block_type == "heading":
            assert block.heading_level is not None
            assert 1 <= block.heading_level <= 6
        else:
            assert block.heading_level is None


def test_title_level_follows_outline_indent_and_is_clamped(tmp_path: Path) -> None:
    """标题段落带大纲缩进时按 `lvl + 1` 算；越界（lvl=8）夹到 6。"""
    path = tmp_path / "deep-title.pptx"
    prs = Presentation()
    for level in (1, 8):
        slide = prs.slides.add_slide(_layout(prs, "Title and Content"))
        slide.shapes.title.text = f"缩进级别 {level}"
        slide.shapes.title.text_frame.paragraphs[0].level = level
    prs.save(str(path))

    headings = parse_pptx(path).headings
    assert [h.heading_level for h in headings] == [2, 6]


# ---------------------------------------------------------------------------
# 正文
# ---------------------------------------------------------------------------


def test_body_text_and_bullets_become_paragraphs(fixture_pptx: Path) -> None:
    """正文（含二级项目符号、非占位符文本框）一律 `paragraph`，内容保持原貌。"""
    doc = parse_pptx(fixture_pptx)
    paragraphs = {b.content_md for b in doc.blocks if b.block_type == "paragraph"}
    assert "路由是把分组从源送到目的的过程。" in paragraphs
    assert "路由器根据路由表决定下一跳。" in paragraphs
    assert "最长前缀匹配" in paragraphs
    assert "补充说明：这张表对比两种协议。" in paragraphs


def test_one_block_per_paragraph(fixture_pptx: Path) -> None:
    """一个段落一块：项目符号列表里每条都是独立的知识单元，不合成一大块。"""
    doc = parse_pptx(fixture_pptx)
    page1_paragraphs = [b.content_md for b in doc.blocks if b.page_no == 1 and not b.is_heading]
    assert page1_paragraphs[:3] == [
        "路由是把分组从源送到目的的过程。",
        "路由器根据路由表决定下一跳。",
        "最长前缀匹配",
    ]


def test_blank_text_box_produces_no_block(tmp_path: Path) -> None:
    """空文本框不产块（否则会变成一个溯不到任何原文的噪声块）。"""
    path = tmp_path / "blank-box.pptx"
    prs = Presentation()
    slide = _blank_slide(prs)
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "有内容"
    prs.save(str(path))

    doc = parse_pptx(path)
    assert [b.content_md for b in doc.blocks] == ["有内容"]


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------


def test_table_becomes_a_markdown_table(fixture_pptx: Path) -> None:
    tables = [b for b in parse_pptx(fixture_pptx).blocks if b.block_type == "table"]
    assert len(tables) == 1
    md = tables[0].content_md
    assert "| 协议 | 作用 | 层次 |" in md
    assert "| --- | --- | --- |" in md
    assert "| TCP | 可靠传输 | 传输层 |" in md
    assert tables[0].page_no == 2


def test_horizontally_merged_cell_is_written_once(fixture_pptx: Path) -> None:
    """★ 横向合并：内容只写合并起点那一格，续格留空 —— 不许复制成 n 份。

    Markdown 没有 colspan，重复写内容会让"跨两列"看起来像两列各自的取值，
    下游按列取值时直接读错。
    """
    table = next(b for b in parse_pptx(fixture_pptx).blocks if b.block_type == "table")
    lines = table.content_md.splitlines()
    merged_lines = [line for line in lines if "跨两列" in line]
    assert len(merged_lines) == 1
    assert merged_lines[0].count("跨两列") == 1
    # 列数必须每行一致，否则多数渲染器会把整张表渲染塌
    assert {line.count("|") for line in lines} == {4}


def test_empty_table_produces_no_block(fixture_pptx: Path) -> None:
    """第 4 张只有一张全空表格。`|  |` + 分隔行 strip() 后仍含 `|`，会躲过空白
    块过滤变成噪声块，必须在表格渲染里就返回空串。"""
    doc = parse_pptx(fixture_pptx)
    assert [b for b in doc.blocks if b.page_no == 4] == []
    assert len([b for b in doc.blocks if b.block_type == "table"]) == 1


# ---------------------------------------------------------------------------
# 备注（notes_slide）
# ---------------------------------------------------------------------------


def test_notes_become_other_blocks_with_a_marker(fixture_pptx: Path) -> None:
    """★ 备注要产出，且必须能与正文区分：`block_type="other"` + `> 备注：` 前缀。"""
    doc = parse_pptx(fixture_pptx)
    notes = [b for b in doc.blocks if b.block_type == "other"]
    assert [b.content_md for b in notes] == [
        "> 备注：备注第一段：先讲清楚“转发”与“路由”的区别。",
        "> 备注：备注第二段：留 2 分钟做课堂提问。",
    ]
    assert all(b.content_md.startswith("> 备注：") for b in notes)
    assert {b.page_no for b in notes} == {1}
    assert all(b.heading_level is None for b in notes)

    # 两层区分都要真的起作用：类型不同 + 前缀不会被正文文本命中
    body_texts = {b.content_md for b in doc.blocks if b.block_type != "other"}
    assert all(n.content_md not in body_texts for n in notes)


def test_notes_follow_their_own_slide(fixture_pptx: Path) -> None:
    """备注紧跟它所属的那一页，人工核对"这页讲了什么"时顺序顺眼。"""
    doc = parse_pptx(fixture_pptx)
    page1_types = [b.block_type for b in doc.blocks if b.page_no == 1]
    assert page1_types == ["heading", "paragraph", "paragraph", "paragraph", "other", "other"]
    assert doc.blocks[6].page_no == 2  # 第 2 张紧接在第 1 张的备注之后


def test_slide_without_notes_produces_no_other_block(tmp_path: Path) -> None:
    """没有备注就一个 `other` 块都不许有 —— **尤其不许为了读备注创建备注页**。

    `slide.notes_slide` 在没有备注页时会**凭空创建一张**，那是写操作；
    解析器只读，必须走 `has_notes_slide` 先探。
    """
    path = tmp_path / "no-notes.pptx"
    prs = Presentation()
    slide = _blank_slide(prs)
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "只有正文"
    prs.save(str(path))

    doc = parse_pptx(path)
    assert doc.blocks and all(b.block_type != "other" for b in doc.blocks)

    # 再打开一次：如果解析时创建过备注页，这里会看到它被**持久化**了
    reopened = Presentation(str(path))
    assert not reopened.slides[0].has_notes_slide


# ---------------------------------------------------------------------------
# 页码 / 行号 / 序号
# ---------------------------------------------------------------------------


def test_page_no_is_1_based_slide_index_on_every_block(fixture_pptx: Path) -> None:
    """★ 第 1 张幻灯片就是第 1 页；空白幻灯片照样占页码（第 4 张仍是 4）。"""
    doc = parse_pptx(fixture_pptx)
    assert all(b.page_no is not None for b in doc.blocks)
    assert {b.page_no for b in doc.blocks} == {1, 2}
    assert doc.page_count == 4
    # 第 3 张是空白页：不产块，但也不许把它"挤掉"让第 4 张变成第 3 页
    assert [b for b in doc.blocks if b.page_no == 3] == []


def test_line_numbers_are_none_for_every_block(fixture_pptx: Path) -> None:
    """PPTX 没有"页内行号"这个坐标系（换行是排版结果，resize 一下就变），
    给一个会漂移的行号比没有行号更糟。"""
    doc = parse_pptx(fixture_pptx)
    assert all(b.line_start is None and b.line_end is None for b in doc.blocks)


def test_seq_is_global_continuous_from_zero(fixture_pptx: Path) -> None:
    """★ seq 是整份材料的全局序号：0,1,2…，跨幻灯片也不重置。"""
    doc = parse_pptx(fixture_pptx)
    assert [seq for seq, _ in doc.numbered()] == list(range(len(doc.blocks)))


def test_block_ids_match_the_canonical_id_helper(fixture_pptx: Path) -> None:
    """★ 锚点里的 block_id 必须与 `app.models.ids.block_id()` **现算**一致（A1-6）。"""
    doc = parse_pptx(fixture_pptx)
    mat_id = material_id(hash_bytes(fixture_pptx.read_bytes()))
    anchors = _ANCHOR_RE.findall(to_markdown(doc, mat_id))
    assert anchors == [block_id(mat_id, seq) for seq in range(len(doc.blocks))]
    assert anchors[0].startswith("blk_")


# ---------------------------------------------------------------------------
# 坏输入：一律中文 ApiError（A1-7 失败隔离）
# ---------------------------------------------------------------------------


def test_empty_file_raises_chinese_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.pptx"
    path.write_bytes(b"")
    with pytest.raises(ApiError) as excinfo:
        parse_pptx(path)
    _assert_chinese_error(excinfo.value, "空文件")


def test_deck_without_any_slide_raises_chinese_error(tmp_path: Path) -> None:
    """0 字节之外，"一张幻灯片都没有"的空 PPT 也要拒绝：它会让材料显示
    "解析成功"却零产物，是最不该静默的那种失败。"""
    path = tmp_path / "no-slides.pptx"
    Presentation().save(str(path))
    with pytest.raises(ApiError) as excinfo:
        parse_pptx(path)
    _assert_chinese_error(excinfo.value, "任何幻灯片")


def test_corrupted_pptx_raises_chinese_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"this is definitely not a pptx")
    with pytest.raises(ApiError) as excinfo:
        parse_pptx(path)
    _assert_chinese_error(excinfo.value, "损坏")


def test_legacy_ppt_or_encrypted_raises_chinese_error(tmp_path: Path) -> None:
    """OLE2 头 = 加密的 .pptx 或旧版 .ppt，处置动作相同（另存为 .pptx）。"""
    path = tmp_path / "legacy.pptx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ApiError) as excinfo:
        parse_pptx(path)
    _assert_chinese_error(excinfo.value, "加密")


# ---------------------------------------------------------------------------
# 入口分派 / 扩展名
# ---------------------------------------------------------------------------


def test_parse_material_dispatches_pptx(fixture_pptx: Path) -> None:
    doc = parse_material(fixture_pptx)
    assert doc.source_type == "pptx"
    assert doc == parse_pptx(fixture_pptx)


def test_uppercase_extension_is_dispatched_too(tmp_path: Path) -> None:
    """用户上传的 `讲义.PPTX` 不该被当成未知格式。"""
    path = make_deck(tmp_path / "讲义.PPTX")
    assert parse_material(path).source_type == "pptx"


def test_legacy_ppt_is_still_rejected_but_points_at_pptx(tmp_path: Path) -> None:
    """旧版 `.ppt` 保持拒绝，但文案必须指路"另存为 .pptx"（用户自己就能解决）。"""
    path = tmp_path / "讲义.ppt"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ApiError) as excinfo:
        parse_material(path)

    err = excinfo.value
    assert err.code == ErrorCode.UNSUPPORTED_FORMAT
    assert err.status_code == 400
    assert _HAN_RE.search(err.message)
    assert ".pptx" in err.message and "另存为" in err.message


# ---------------------------------------------------------------------------
# 不确定处
# ---------------------------------------------------------------------------


def test_deck_without_title_placeholder_says_so(tmp_path: Path) -> None:
    """没有标题占位符 → 骨架推不出来，必须给一条能展示的中文说明。"""
    path = tmp_path / "no-title.pptx"
    prs = Presentation()
    slide = _blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "只有正文"
    prs.save(str(path))

    doc = parse_pptx(path)
    assert doc.blocks and not doc.headings
    assert [n.kind for n in doc.uncertain_notes] == ["ambiguous_structure"]
    assert "标题" in doc.uncertain_notes[0].message


def test_deck_with_title_has_no_structure_note(fixture_pptx: Path) -> None:
    assert parse_pptx(fixture_pptx).uncertain_notes == []


# ---------------------------------------------------------------------------
# 可复现（A2-7）与落库
# ---------------------------------------------------------------------------


def test_same_pptx_parses_identically_twice(fixture_pptx: Path) -> None:
    assert parse_pptx(fixture_pptx) == parse_pptx(fixture_pptx)
    assert parse_material(fixture_pptx) == parse_material(fixture_pptx)


def test_persist_blocks_keeps_slide_numbers(fixture_pptx: Path, session: Session) -> None:
    """落库：幻灯片号进 `page_no`、音频字段不写（D-08）、seq 与块锚点同源。"""
    session.add(
        Material(
            id=MAT_ID,
            filename="讲义.pptx",
            file_hash=hash_bytes(b"pptx fixture"),
            stored_path="uploads/讲义.pptx",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            size_bytes=fixture_pptx.stat().st_size,
            source_type="pptx",
            parse_method="text_extract",
        )
    )
    session.flush()

    doc = parse_pptx(fixture_pptx)
    ids = persist_blocks(session, MAT_ID, doc)

    rows = (
        session.execute(
            select(MaterialBlock)
            .where(MaterialBlock.material_id == MAT_ID)
            .order_by(MaterialBlock.seq)
        )
        .scalars()
        .all()
    )
    assert [row.seq for row in rows] == list(range(len(doc.blocks)))
    assert ids == [block_id(MAT_ID, seq) for seq in range(len(doc.blocks))]
    assert [row.page_no for row in rows] == [b.page_no for b in doc.blocks]
    for row in rows:
        assert row.line_start is None and row.line_end is None
        assert row.ts_start_ms is None and row.ts_end_ms is None
        assert row.ocr_confidence is None
