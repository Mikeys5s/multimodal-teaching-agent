"""QA 独立验证 · 场景 12–13：锚点与可复现（A1-6 / A2-7）+ 落库一致性。

覆盖点：
  · 场景 12：PDF 每个块都有 `page_no` 且落在 `[1, page_count]`；
    块锚点 `<!-- block: blk_... -->` 数量与块数一致，且 ID **现算**后逐一相等
  · 场景 13：同一输入连续解析两次，块序列（id/seq/page_no/content_md）与
    `to_markdown()` 输出**逐字节相同**
  · 额外：`persist_blocks` 落库后，库里的 id/seq/page_no 与 Markdown 锚点同源，
    音频字段（D-08）与 OCR 置信度必须为 NULL

反向验证（防"假绿"）：本文件专门断言 conftest 的引擎**确实开了外键**
（往不存在的 material_id 落块必须 IntegrityError）。如果哪天 conftest 换成
自建引擎却忘了 `apply_sqlite_pragmas`，这条会先红。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import docx
import pymupdf
import pytest
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.ids import block_id, hash_bytes, material_id
from app.models.material import Material, MaterialBlock
from app.parse import parse_docx, parse_pdf, to_markdown
from app.parse.persist import persist_blocks

CJK_FONT = "china-s"

#: 材料 ID 现算：`mat_<sha256 前 8 位>`。块 ID 必须由它 + seq 现算得出。
MAT_ID = material_id(hash_bytes(b"qa-anchor-fixture"))
FILE_HASH = hash_bytes(b"qa-anchor-fixture")

_BLOCK_ANCHOR_RE = re.compile(r"<!-- block: (blk_[0-9a-f]+_\d{5}) -->")
_PAGE_ANCHOR_RE = re.compile(r"<!-- page: (\d+) -->")


def _w(page: pymupdf.Page, x: float, y: float, text: str, size: float = 10) -> None:
    page.insert_text((x, y), text, fontname=CJK_FONT, fontsize=size)


def make_three_page_pdf(path: Path) -> Path:
    """三页：每页都有标题 + 正文，用来验证页码锚点与块锚点。"""
    doc = pymupdf.open()

    p1 = doc.new_page()
    _w(p1, 72, 72, "第1章 网络层", 18)
    _w(p1, 72, 110, "1.1 路由基础", 13)
    _w(p1, 72, 140, "路由是把分组从源送到目的的过程。", 10)
    _w(p1, 72, 170, "路由表里保存着下一跳地址。", 10)

    p2 = doc.new_page()
    _w(p2, 72, 110, "1.2 转发与路由选择", 13)
    _w(p2, 72, 140, "转发是本地动作。", 10)

    p3 = doc.new_page()
    _w(p3, 72, 110, "第2章 应用层", 18)
    _w(p3, 72, 140, "应用层协议定义了报文的格式与次序。", 10)

    doc.save(str(path))
    doc.close()
    return path


def make_two_page_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("第1章 网络层", level=1)
    document.add_paragraph("路由是把分组从源送到目的的过程。")
    document.add_table(rows=2, cols=2)
    document.tables[0].cell(0, 0).text = "协议"
    document.tables[0].cell(0, 1).text = "作用"
    document.tables[0].cell(1, 0).text = "IP"
    document.tables[0].cell(1, 1).text = "寻址"
    document.add_paragraph("表格后面还有正文。")
    document.save(str(path))
    return path


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    return make_three_page_pdf(tmp_path / "anchors.pdf")


def _signature(doc) -> list[tuple[str, int, int | None, str]]:
    """`(block_id, seq, page_no, content_md)` —— 场景 13 要逐条比对的那个序列。"""
    return [
        (block_id(MAT_ID, seq), seq, block.page_no, block.content_md)
        for seq, block in doc.numbered()
    ]


def _add_material(session: Session, name: str, source_type: str, parse_method: str | None) -> None:
    session.add(
        Material(
            id=MAT_ID,
            filename=name,
            file_hash=FILE_HASH,
            stored_path=f"uploads/{name}",
            mime_type="application/octet-stream",
            size_bytes=1,
            source_type=source_type,
            parse_method=parse_method,
        )
    )
    session.flush()


# ---------------------------------------------------------------------------
# 场景 12：锚点
# ---------------------------------------------------------------------------


def test_every_pdf_block_has_page_no_within_range(pdf_path: Path) -> None:
    """★ 场景 12：每个块都有 `page_no`，且落在 `[1, page_count]`。"""
    doc = parse_pdf(pdf_path)
    assert doc.page_count == 3
    assert doc.blocks
    for block in doc.blocks:
        assert block.page_no is not None, "文本层 PDF 的块必须有页码"
        assert 1 <= block.page_no <= doc.page_count


def test_block_page_numbers_are_1_based_and_every_page_has_a_block(pdf_path: Path) -> None:
    """★ 场景 12：页码是 1-based（不许出现 0），且每页都有块、页码不回跳。

    断言的是合成 3 页 PDF 里**块自身的 `page_no`** 集合（== {1,2,3}），
    与 `to_markdown()` 的页锚点无关 —— 页锚点由 `test_page_anchors_*` 覆盖。
    """
    doc = parse_pdf(pdf_path)
    pages = [b.page_no for b in doc.blocks]
    assert min(pages) == 1
    assert set(pages) == {1, 2, 3}
    assert pages == sorted(pages), "块顺序必须与文档顺序一致，页码不能回跳"


def test_docx_blocks_have_no_page_no(tmp_path: Path) -> None:
    """DOCX 无页概念 —— 不能编造页码（`page_no` 全为 None）。"""
    doc = parse_docx(make_two_page_docx(tmp_path / "anchors.docx"))
    assert all(b.page_no is None for b in doc.blocks)


def test_block_anchor_count_matches_block_count(pdf_path: Path) -> None:
    """★ 场景 12：块锚点数量与该材料的块数**严格相等**（多一个少一个都是溯源错位）。"""
    doc = parse_pdf(pdf_path)
    md = to_markdown(doc, MAT_ID)
    anchors = _BLOCK_ANCHOR_RE.findall(md)
    assert len(anchors) == len(doc.blocks)
    assert len(anchors) == len(set(anchors)), "块锚点出现重复"


def test_block_anchor_ids_equal_recomputed_block_ids(pdf_path: Path) -> None:
    """★ 场景 12：锚点里的 ID 必须等于 `ids.block_id(mat_id, seq)` 现算的结果，且顺序一致。"""
    doc = parse_pdf(pdf_path)
    md = to_markdown(doc, MAT_ID)
    anchors = _BLOCK_ANCHOR_RE.findall(md)
    expected = [block_id(MAT_ID, seq) for seq in range(len(doc.blocks))]
    assert anchors == expected
    for anchor in anchors:
        assert re.fullmatch(r"blk_[0-9a-f]{8}_\d{5}", anchor), f"锚点格式不对：{anchor}"


def test_every_block_appears_in_markdown_in_order(pdf_path: Path) -> None:
    """每个块的内容都要在 Markdown 里出现，且锚点与内容的相对顺序正确。"""
    doc = parse_pdf(pdf_path)
    md = to_markdown(doc, MAT_ID)
    cursor = -1
    # 锚点与块必须一一对应：长度不等说明有块没进 Markdown，或反之
    for anchor, block in zip(_BLOCK_ANCHOR_RE.findall(md), doc.blocks, strict=True):
        idx = md.index(anchor)
        assert idx > cursor
        cursor = idx
        assert block.content_md.strip() in md


def test_page_anchor_is_event_based(pdf_path: Path) -> None:
    """页码锚点是事件式的：同页不重复插，跨页才出现。"""
    doc = parse_pdf(pdf_path)
    md = to_markdown(doc, MAT_ID)
    pages = _PAGE_ANCHOR_RE.findall(md)
    assert pages == ["1", "2", "3"], f"页码锚点数量不对：{pages}"
    assert md.count("<!-- page: 1 -->") == 1


def test_page_anchor_precedes_the_first_block_of_that_page(pdf_path: Path) -> None:
    """页码锚点必须落在该页第一个块锚点之前（否则页码会被归到上一页）。"""
    doc = parse_pdf(pdf_path)
    md = to_markdown(doc, MAT_ID)
    first_block_of_page: dict[int, str] = {}
    for seq, block in doc.numbered():
        first_block_of_page.setdefault(block.page_no, block_id(MAT_ID, seq))
    for page_no, anchor in first_block_of_page.items():
        page_pos = md.index(f"<!-- page: {page_no} -->")
        block_pos = md.index(f"<!-- block: {anchor} -->")
        assert page_pos < block_pos


def test_docx_markdown_has_block_anchors_but_no_page_anchors(tmp_path: Path) -> None:
    """★ DOCX：块锚点照旧，页码锚点一个都不能有（否则凭空冒出一堆 `page:`）。"""
    doc = parse_docx(make_two_page_docx(tmp_path / "anchors.docx"))
    md = to_markdown(doc, MAT_ID)
    assert len(_BLOCK_ANCHOR_RE.findall(md)) == len(doc.blocks)
    assert _PAGE_ANCHOR_RE.findall(md) == []
    assert "<!-- page:" not in md


def test_start_seq_offsets_anchor_ids(tmp_path: Path) -> None:
    """切片场景：`start_seq` 必须体现在锚点 ID 上（重新从 0 编号会让锚点指错块）。"""
    doc = parse_pdf(make_three_page_pdf(tmp_path / "anchors.pdf"))
    md = to_markdown(doc.blocks[:2], MAT_ID, start_seq=5)
    assert _BLOCK_ANCHOR_RE.findall(md) == [block_id(MAT_ID, 5), block_id(MAT_ID, 6)]


# ---------------------------------------------------------------------------
# 场景 13：可复现
# ---------------------------------------------------------------------------


def test_two_parses_produce_identical_block_sequence(pdf_path: Path) -> None:
    """★ 场景 13：两次解析的 `(block_id, seq, page_no, content_md)` 序列完全相同。"""
    first = parse_pdf(pdf_path)
    second = parse_pdf(pdf_path)
    assert _signature(first) == _signature(second)
    assert len(_signature(first)) == len(first.blocks)


def test_two_parses_produce_byte_identical_markdown(pdf_path: Path) -> None:
    """★ 场景 13：`to_markdown()` 输出**逐字节相同**（不是"内容差不多"）。"""
    first_md = to_markdown(parse_pdf(pdf_path), MAT_ID)
    second_md = to_markdown(parse_pdf(pdf_path), MAT_ID)
    assert first_md.encode("utf-8") == second_md.encode("utf-8")


def test_three_parses_are_stable(pdf_path: Path) -> None:
    """第三次、第四次也不能漂移（防止"第一次初始化了什么"这类状态问题）。"""
    baseline_md = to_markdown(parse_pdf(pdf_path), MAT_ID)
    baseline_sig = _signature(parse_pdf(pdf_path))
    for _ in range(3):
        assert to_markdown(parse_pdf(pdf_path), MAT_ID) == baseline_md
        assert _signature(parse_pdf(pdf_path)) == baseline_sig


def test_docx_parses_identically_twice(tmp_path: Path) -> None:
    """DOCX 同样要可复现（顺序、表格渲染、锚点都不能漂）。"""
    path = make_two_page_docx(tmp_path / "anchors.docx")
    assert parse_docx(path) == parse_docx(path)
    assert to_markdown(parse_docx(path), MAT_ID) == to_markdown(parse_docx(path), MAT_ID)


def test_reparse_after_a_different_file_does_not_change_result(pdf_path: Path, tmp_path: Path) -> None:
    """解析顺序不能影响结果：先解析别的文件，再回来解析，结果必须一样。"""
    baseline = _signature(parse_pdf(pdf_path))
    parse_pdf(make_three_page_pdf(tmp_path / "other.pdf"))
    parse_docx(make_two_page_docx(tmp_path / "other.docx"))
    assert _signature(parse_pdf(pdf_path)) == baseline


# ---------------------------------------------------------------------------
# 落库一致性（锚点 ↔ 数据库行）
# ---------------------------------------------------------------------------


def test_persisted_rows_match_parse_product(pdf_path: Path, session: Session) -> None:
    """★ 库里的 id / seq / page_no / content_md 必须与解析产物一一对应。"""
    _add_material(session, "anchors.pdf", "pdf_text", "text_extract")
    doc = parse_pdf(pdf_path)
    ids = persist_blocks(session, MAT_ID, doc)

    rows = (
        session.execute(
            select(MaterialBlock).where(MaterialBlock.material_id == MAT_ID).order_by(MaterialBlock.seq)
        )
        .scalars()
        .all()
    )

    assert [row.seq for row in rows] == list(range(len(doc.blocks)))
    assert [row.id for row in rows] == ids
    assert ids == [block_id(MAT_ID, seq) for seq in range(len(doc.blocks))]

    # 前面已断言两者长度相等，这里用 strict=True 把它变成循环本身的前置检查
    for row, block in zip(rows, doc.blocks, strict=True):
        assert row.id == block_id(MAT_ID, row.seq)
        assert row.page_no == block.page_no
        assert row.content_md == block.content_md
        assert row.block_type == block.block_type
        assert row.heading_level == block.heading_level
        assert row.line_start == block.line_start and row.line_end == block.line_end
        if block.bbox is None:
            assert row.bbox is None
        else:
            assert json.loads(row.bbox) == [round(float(v), 2) for v in block.bbox]


def test_persisted_rows_ids_equal_markdown_anchors(pdf_path: Path, session: Session) -> None:
    """★ 场景 12 的地基：库里的主键 == Markdown 锚点里的 ID（同源，不可能不一致）。"""
    _add_material(session, "anchors.pdf", "pdf_text", "text_extract")
    doc = parse_pdf(pdf_path)
    ids = persist_blocks(session, MAT_ID, doc)
    anchors = _BLOCK_ANCHOR_RE.findall(to_markdown(doc, MAT_ID))
    assert ids == anchors


def test_audio_and_ocr_fields_stay_null(pdf_path: Path, session: Session) -> None:
    """★ D-08：`ts_start_ms` / `ts_end_ms` 保留不启用；本批次没有 OCR，置信度也不许写。"""
    _add_material(session, "anchors.pdf", "pdf_text", "text_extract")
    persist_blocks(session, MAT_ID, parse_pdf(pdf_path))

    rows = (
        session.execute(select(MaterialBlock).where(MaterialBlock.material_id == MAT_ID))
        .scalars()
        .all()
    )
    assert rows
    for row in rows:
        assert row.ts_start_ms is None
        assert row.ts_end_ms is None
        assert row.ocr_confidence is None


def test_persist_writes_nothing_for_a_scan_document(tmp_path: Path, session: Session) -> None:
    """扫描版不产出块 → 库里也不该有任何行（否则就是伪造内容）。"""
    doc_page = pymupdf.open()
    page = doc_page.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 200))
    pix.clear_with(180)
    page.insert_image(pymupdf.Rect(40, 40, 340, 240), pixmap=pix)
    path = tmp_path / "scan.pdf"
    doc_page.save(str(path))
    doc_page.close()

    _add_material(session, "scan.pdf", "pdf_scan", None)
    doc = parse_pdf(path)
    ids = persist_blocks(session, MAT_ID, doc)
    assert ids == []
    count = session.execute(
        select(func.count()).select_from(MaterialBlock).where(MaterialBlock.material_id == MAT_ID)
    ).scalar_one()
    assert count == 0


def test_persist_does_not_commit_by_itself(pdf_path: Path, session: Session) -> None:
    """`persist_blocks` 不 commit —— 事务边界必须留给调用方（否则会出现"状态 done 但块没落"）。"""
    _add_material(session, "anchors.pdf", "pdf_text", "text_extract")
    doc = parse_pdf(pdf_path)
    persist_blocks(session, MAT_ID, doc)
    # 同一 session 里能查到（已 flush）
    assert session.execute(
        select(func.count()).select_from(MaterialBlock).where(MaterialBlock.material_id == MAT_ID)
    ).scalar_one() == len(doc.blocks)
    assert session.in_transaction()


def test_foreign_key_is_actually_enforced(session: Session) -> None:
    """★ 防"假绿"：conftest 的引擎必须真的开了外键。

    如果测试引擎漏了 `apply_sqlite_pragmas`（SQLite 默认 `foreign_keys=OFF`），
    下面这条**不会**报错 —— 那么本文件所有落库断言都只是在验证"写进去了"，
    却验证不了"写对了地方"。
    """
    session.add(
        MaterialBlock(
            id="blk_deadbeef_00000",
            material_id="mat_doesnotexist",
            seq=0,
            block_type="paragraph",
            content_md="孤块",
        )
    )
    with pytest.raises(IntegrityError):
        session.flush()
