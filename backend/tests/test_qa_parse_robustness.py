"""QA 独立验证 · 场景 1–5：输入健壮性与失败隔离（A1-7）。

本文件只做黑盒验证：一律通过稳定入口 `parse_material(path)` 调用，不直接
触碰解析器内部函数 —— 开发 agent 的 `test_parse_*.py` 已经覆盖了
`parse_pdf` / `parse_docx` 的内部细节，这里要验的是"线上真正被调用的那个入口"
在坏输入下是否给出**能直接展示给用户的中文**，以及一个文件失败后
**下一个文件是否还能正常解析**（失败隔离）。

断言口径（刻意从严）：
  · 异常类型必须是 `ApiError`（业务异常），不能是 `FileNotFoundError` / `OSError` /
    pymupdf 的 `FileDataError` —— 后者会绕过全局异常处理器变成 500 + 英文堆栈；
  · `message` 必须含中文，且**不许**出现 `Traceback` / `Error` / `Exception` /
    `Errno` / 底层英文原话（"cannot open"、"broken document" 等）；
  · 不支持格式的 `message` 必须**点名具体格式**（"PPTX"、"旧版 .doc"、"音频"），
    并给出替代做法（转 PDF / DOCX）。

⚠️ **图片不再属于"未接格式"**（D3 之后）：`.png` / `.jpg` / `.jpeg` / `.webp` /
`.tif` / `.tiff` 走 `app/parse/ocr.py` 的 OCR 通道（需要 `[ocr]` 可选依赖，见
`test_parse_ocr.py`）。所以本文件把它们从"未接格式"名单里摘掉了 —— 它们现在
只在**文件本身坏了**时才报错，那条中文文案由 `test_parse_ocr.py` 守。

夹具全部程序化生成，不依赖仓库外的素材文件。
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pymupdf
import pytest

from app.core.errors import ApiError, ErrorCode
from app.parse import parse_material

CJK_FONT = "china-s"
_HAN_RE = re.compile(r"[\u4e00-\u9fff]")

#: 底层库/解释器可能泄漏到 message 里的英文原话（出现即判定为泄漏）。
_RAW_EXCEPTION_TOKENS = (
    "Traceback",
    "Exception",
    "Errno",
    "cannot open",
    "broken document",
    "BadZipFile",
    "no such file",
    "No such file",
)
_RAW_EXCEPTION_RE = re.compile(r"\b(?:Error|KeyError|ValueError|RuntimeError)\b")


# ---------------------------------------------------------------------------
# 夹具
# ---------------------------------------------------------------------------


def _page_with_text(page: pymupdf.Page, text: str, size: float = 12) -> None:
    page.insert_text((72, 100), text, fontname=CJK_FONT, fontsize=size)


def make_ok_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    _page_with_text(page, "第1章 网络层", 18)
    page.insert_text((72, 140), "路由是分组从源到目的的过程。", fontname=CJK_FONT, fontsize=10)
    doc.save(str(path))
    doc.close()
    return path


def make_ok_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("第1章 网络层", level=1)
    document.add_paragraph("路由是分组从源到目的的过程。")
    document.save(str(path))
    return path


def make_encrypted_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    _page_with_text(page, "这份文件被加密了")
    doc.save(
        str(path),
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="user-secret",
    )
    doc.close()
    return path


def _assert_clean_chinese_message(err: ApiError) -> None:
    """报错文案必须是"能直接给用户看的中文"，不是技术描述。"""
    message = err.message
    assert isinstance(message, str) and message
    assert _HAN_RE.search(message), f"报错文案里没有中文：{message!r}"
    for token in _RAW_EXCEPTION_TOKENS:
        assert token not in message, f"报错文案泄漏了底层英文 {token!r}：{message!r}"
    assert not _RAW_EXCEPTION_RE.search(message), f"报错文案泄漏了异常类名：{message!r}"
    # 面向用户的文案要给出下一步动作
    assert "请" in message, f"报错文案没有给出下一步动作：{message!r}"


# ---------------------------------------------------------------------------
# 场景 1：文件不存在
# ---------------------------------------------------------------------------


def test_missing_file_raises_api_error_not_file_not_found(tmp_path: Path) -> None:
    """★ 场景 1：不存在的路径必须是业务异常 `ApiError`，不是 `FileNotFoundError`。

    `FileNotFoundError` 会绕过全局异常处理器（只处理 ApiError / HTTPException /
    兜底 Exception），用户拿到的会是 500 与英文堆栈。
    """
    missing = tmp_path / "根本不存在.pdf"
    with pytest.raises(ApiError) as excinfo:
        parse_material(missing)

    err = excinfo.value
    assert not isinstance(err, FileNotFoundError)
    assert not isinstance(err, OSError)
    assert not issubclass(ApiError, OSError)
    assert err.code == ErrorCode.NOT_FOUND
    assert err.status_code == 404
    _assert_clean_chinese_message(err)


def test_missing_directory_is_treated_as_missing_file(tmp_path: Path) -> None:
    """传进来一个目录（而不是文件）也要走同一条中文报错，不能崩栈。"""
    with pytest.raises(ApiError) as excinfo:
        parse_material(tmp_path)
    assert excinfo.value.code == ErrorCode.NOT_FOUND
    _assert_clean_chinese_message(excinfo.value)


# ---------------------------------------------------------------------------
# 场景 2：0 字节文件
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", ["empty.pdf", "empty.docx"])
def test_zero_byte_file_raises_clear_chinese_error(tmp_path: Path, name: str) -> None:
    """★ 场景 2：0 字节文件当 PDF / DOCX 传进来 → 明确中文报错，不崩栈。"""
    target = tmp_path / name
    target.write_bytes(b"")

    with pytest.raises(ApiError) as excinfo:
        parse_material(target)

    err = excinfo.value
    assert err.code == ErrorCode.INVALID_PARAM
    assert err.status_code == 400
    _assert_clean_chinese_message(err)


def test_zero_byte_pdf_says_it_is_an_empty_file(tmp_path: Path) -> None:
    """PDF 分支的文案要直接说"空文件"，用户才知道该重新导出。"""
    target = tmp_path / "empty.pdf"
    target.write_bytes(b"")
    with pytest.raises(ApiError) as excinfo:
        parse_material(target)
    assert "空文件" in excinfo.value.message


def test_zero_byte_docx_message_is_actionable(tmp_path: Path) -> None:
    """DOCX 分支的文案要给出可执行动作。

    两个分支措辞不同：0 字节的 .docx 报"已损坏或不是有效的 .docx"，0 字节的 .pdf
    报"空文件"。两者都是明确、可展示的中文，也都给出了可执行动作，故不做统一。
    """
    target = tmp_path / "empty.docx"
    target.write_bytes(b"")
    with pytest.raises(ApiError) as excinfo:
        parse_material(target)
    message = excinfo.value.message
    assert "重新导出" in message
    assert ".docx" in message


# ---------------------------------------------------------------------------
# 场景 3：加密 PDF / 随机字节伪装成 .pdf
# ---------------------------------------------------------------------------


def test_encrypted_pdf_raises_chinese_error_through_entry_point(tmp_path: Path) -> None:
    """★ 场景 3：加密 PDF 走 `parse_material` 也要是中文业务异常。"""
    path = make_encrypted_pdf(tmp_path / "secret.pdf")

    with pytest.raises(ApiError) as excinfo:
        parse_material(path)

    err = excinfo.value
    assert err.code == ErrorCode.INVALID_PARAM
    _assert_clean_chinese_message(err)
    assert "加密" in err.message
    assert "密码" in err.message


def test_random_bytes_named_pdf_raises_chinese_error(tmp_path: Path) -> None:
    """★ 场景 3：随机字节伪装成 .pdf（连 %PDF 头都没有）→ 中文报错。"""
    path = tmp_path / "fake.pdf"
    path.write_bytes(bytes(range(256)) * 8)

    with pytest.raises(ApiError) as excinfo:
        parse_material(path)

    err = excinfo.value
    assert err.code == ErrorCode.INVALID_PARAM
    _assert_clean_chinese_message(err)
    assert "损坏" in err.message or "不是有效的" in err.message


def test_pdf_header_only_truncated_file_raises_chinese_error(tmp_path: Path) -> None:
    """只有 PDF 头、没有正文的截断文件 —— 同样不能崩栈。"""
    path = tmp_path / "truncated.pdf"
    path.write_bytes(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")

    with pytest.raises(ApiError) as excinfo:
        parse_material(path)
    _assert_clean_chinese_message(excinfo.value)


def test_random_bytes_named_docx_raises_chinese_error(tmp_path: Path) -> None:
    """随机字节伪装成 .docx → 中文报错（python-docx 的异常类型不稳定，最容易漏网）。"""
    path = tmp_path / "fake.docx"
    path.write_bytes(bytes(range(256)) * 8)

    with pytest.raises(ApiError) as excinfo:
        parse_material(path)
    _assert_clean_chinese_message(excinfo.value)


def test_no_english_stacktrace_leaks_for_any_broken_input(tmp_path: Path) -> None:
    """把几种坏输入串起来，逐条断言文案里没有英文堆栈痕迹。"""
    broken: list[Path] = []
    for name, data in (
        ("a.pdf", b""),
        ("b.pdf", b"not a pdf at all"),
        ("c.docx", b"not a docx at all"),
        ("d.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32),
    ):
        target = tmp_path / name
        target.write_bytes(data)
        broken.append(target)

    for target in broken:
        with pytest.raises(ApiError) as excinfo:
            parse_material(target)
        _assert_clean_chinese_message(excinfo.value)


# ---------------------------------------------------------------------------
# 场景 4：认出但本批次未接的格式 —— 必须点名具体格式
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "keyword"),
    [
        ("讲义.pptx", "PPTX"),
        ("讲义.ppt", "PPT"),
        ("讲义.doc", "旧版 .doc"),
        ("录音.mp3", "音频"),
        ("录音.wav", "音频"),
        ("录音.m4a", "音频"),
    ],
)
def test_pending_format_names_the_concrete_format(tmp_path: Path, filename: str, keyword: str) -> None:
    """★ 场景 4：`.pptx` / 音频 → `UNSUPPORTED_FORMAT` 且 message 点名格式。

    笼统的"不支持"会让用户反复重试；点名"PPTX"并给出替代做法才有用。
    """
    target = tmp_path / filename
    target.write_bytes(b"placeholder bytes")

    with pytest.raises(ApiError) as excinfo:
        parse_material(target)

    err = excinfo.value
    assert err.code == ErrorCode.UNSUPPORTED_FORMAT
    assert err.status_code == 400
    assert keyword in err.message, f"文案没点名具体格式 {keyword!r}：{err.message!r}"
    _assert_clean_chinese_message(err)
    # 必须给出替代做法，否则用户无处可去
    assert "PDF" in err.message and "DOCX" in err.message


def test_unknown_extension_is_echoed_back(tmp_path: Path) -> None:
    """完全没见过的扩展名也要把扩展名回显给用户。"""
    target = tmp_path / "神秘文件.xyz"
    target.write_bytes(b"x")

    with pytest.raises(ApiError) as excinfo:
        parse_material(target)
    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FORMAT
    assert ".xyz" in excinfo.value.message
    _assert_clean_chinese_message(excinfo.value)


def test_file_without_extension_is_reported_as_such(tmp_path: Path) -> None:
    """无扩展名不能说成"暂不支持 格式"（会留下一个空格看起来像 bug）。"""
    target = tmp_path / "noextension"
    target.write_bytes(b"x")

    with pytest.raises(ApiError) as excinfo:
        parse_material(target)
    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FORMAT
    assert "无扩展名" in excinfo.value.message
    _assert_clean_chinese_message(excinfo.value)


def test_unsupported_format_is_decided_by_extension_not_by_content(tmp_path: Path) -> None:
    """扩展名分派必须发生在读文件内容之前 —— 内容是垃圾也不该 500。"""
    target = tmp_path / "空壳.pptx"
    target.write_bytes(b"\x00" * 4)
    with pytest.raises(ApiError) as excinfo:
        parse_material(target)
    assert excinfo.value.code == ErrorCode.UNSUPPORTED_FORMAT


# ---------------------------------------------------------------------------
# 场景 5：失败隔离
# ---------------------------------------------------------------------------


def test_failures_do_not_pollute_following_parses(
    ocr_available_stub: None, tmp_path: Path
) -> None:
    """★ 场景 5：先制造一串失败，再解析正常文件必须仍然成功。

    解析链路的调用方是"逐个文件调用、捕获 ApiError 记到 materials.error_message"，
    所以单文件失败**绝对不能**污染全局状态（模块级缓存、pymupdf 全局文档等）。

    `photo.png` 是坏图（`b"\\x89PNG"`）→ 期望 `UNSUPPORTED_FORMAT`，所以这里挂
    `ocr_available_stub`：把 OCR 固定成"可用"才不会在 `ensure_available()` 就被拦成
    "没装 OCR 依赖"，同时不触发一次真的 paddleocr import（省 ~3 s，也不起引擎）。
    断言一条都没放松 —— 坏图必须在**读图**阶段报"格式不支持"。
    """
    good_pdf = make_ok_pdf(tmp_path / "good.pdf")
    good_docx = make_ok_docx(tmp_path / "good.docx")

    bad_inputs: list[Path] = [
        tmp_path / "missing.pdf",
        tmp_path / "empty.pdf",
        tmp_path / "garbage.pdf",
        tmp_path / "deck.pptx",
        tmp_path / "photo.png",
    ]
    (tmp_path / "empty.pdf").write_bytes(b"")
    (tmp_path / "garbage.pdf").write_bytes(b"garbage not a pdf")
    (tmp_path / "deck.pptx").write_bytes(b"PK")
    (tmp_path / "photo.png").write_bytes(b"\x89PNG")

    seen_codes = []
    for bad in bad_inputs:
        with pytest.raises(ApiError) as excinfo:
            parse_material(bad)
        seen_codes.append(excinfo.value.code)
    assert seen_codes == [
        ErrorCode.NOT_FOUND,
        ErrorCode.INVALID_PARAM,
        ErrorCode.INVALID_PARAM,
        ErrorCode.UNSUPPORTED_FORMAT,
        ErrorCode.UNSUPPORTED_FORMAT,
    ]

    # 失败之后，正常文件必须照常解析
    pdf_doc = parse_material(good_pdf)
    assert pdf_doc.source_type == "pdf_text"
    assert pdf_doc.blocks

    docx_doc = parse_material(good_docx)
    assert docx_doc.source_type == "docx"
    assert docx_doc.blocks

    # 再解析一次同一份 PDF，结果与失败前完全一致（无跨调用状态）
    assert parse_material(good_pdf) == pdf_doc


def test_same_missing_file_fails_identically_every_time(tmp_path: Path) -> None:
    """重复失败也要给出同一条文案（不夹带上一次的状态）。"""
    missing = tmp_path / "nope.docx"
    messages = set()
    for _ in range(3):
        with pytest.raises(ApiError) as excinfo:
            parse_material(missing)
        messages.add((excinfo.value.code, excinfo.value.message))
    assert len(messages) == 1


def test_extension_matching_is_case_insensitive(tmp_path: Path) -> None:
    """用户上传的 `报告.PDF` 不该被当成未知格式。"""
    upper = make_ok_pdf(tmp_path / "报告.PDF")
    assert parse_material(upper).source_type == "pdf_text"

    upper_docx = make_ok_docx(tmp_path / "讲义.DOCX")
    assert parse_material(upper_docx).source_type == "docx"
