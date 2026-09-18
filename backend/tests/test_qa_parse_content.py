"""QA 独立验证 · 场景 6–8、11：内容正确性（分段 / 拼接 / 空白页 / DOCX）。

覆盖点：
  · 场景 6：**段间空行明显**的多段不能被合成一块；同一段折行**应该**合并
  · 场景 7：中英混排、CJK 与拉丁/数字相邻时的拼接，空格有没有多插 / 少插
  · 场景 8：空页 / 只有空白的页不能产出 `content_md` 为空白字符的块
  · 场景 11：DOCX 的标题层级、嵌套表格、文档顺序、`page_no` 必须为 None

夹具全部程序化生成。行距刻意取"单倍行距"（字号 10 → 相邻基线 12）与
"明显段间距"（≥ 40）两种，用它们区分"折行"和"新段落"。
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pymupdf
import pytest

from app.models._common import BLOCK_TYPES
from app.parse import parse_docx, parse_pdf

CJK_FONT = "china-s"
_HAN_RE = re.compile(r"[\u4e00-\u9fff]")

#: 正文 10pt 时的相邻基线差（≈ 单倍行距）
LINE_STEP = 12.0
#: 明显的段间距
PARAGRAPH_GAP = 40.0


def _w(page: pymupdf.Page, x: float, y: float, text: str, size: float = 10) -> None:
    page.insert_text((x, y), text, fontname=CJK_FONT, fontsize=size)


def _one_page_pdf(path: Path, lines: list[tuple[float, float, str] | tuple[float, float, str, float]]) -> Path:
    """`lines` = [(x, y, text)] 或 [(x, y, text, size)]，便于精确控制行距与缩进。"""
    doc = pymupdf.open()
    page = doc.new_page()
    for item in lines:
        if len(item) == 4:
            _w(page, item[0], item[1], item[2], item[3])
        else:
            _w(page, item[0], item[1], item[2])
    doc.save(str(path))
    doc.close()
    return path


def _paragraphs(doc) -> list[str]:
    return [b.content_md for b in doc.blocks if b.block_type == "paragraph"]


# ---------------------------------------------------------------------------
# 场景 6：分段 vs 折行
# ---------------------------------------------------------------------------


def test_blank_line_separated_paragraphs_stay_separate(tmp_path: Path) -> None:
    """★ 场景 6：同一页三段落，段间空行明显 → 必须是三块，绝不能合成一块。

    合成一块会让"这段话属于哪个知识点"彻底含糊，而且 `source_quote` 会引到
    另外两段的内容上。
    """
    path = _one_page_pdf(
        tmp_path / "three_paragraphs.pdf",
        [
            (72, 100, "第一段的内容，讲的是网络层的基本概念。"),
            (72, 100 + PARAGRAPH_GAP, "第二段的内容，和第一段之间有明显空行。"),
            (72, 100 + 2 * PARAGRAPH_GAP, "第三段的内容，同样隔了很远。"),
        ],
    )

    doc = parse_pdf(path)
    assert _paragraphs(doc) == [
        "第一段的内容，讲的是网络层的基本概念。",
        "第二段的内容，和第一段之间有明显空行。",
        "第三段的内容，同样隔了很远。",
    ]
    for block in doc.blocks:
        assert block.page_no == 1
        # 每段都只占一行 —— 顺序与行号都必须对得上
        assert block.line_start == block.line_end
    assert [b.line_start for b in doc.blocks] == [1, 2, 3]


def test_five_blank_line_separated_paragraphs_stay_separate(tmp_path: Path) -> None:
    """段数再多一点，看会不会出现"相邻两块被吃掉"的边界问题。"""
    lines = [(72, 100 + i * PARAGRAPH_GAP, f"第{i + 1}段的内容。") for i in range(5)]
    doc = parse_pdf(_one_page_pdf(tmp_path / "five.pdf", lines))
    assert _paragraphs(doc) == [f"第{i + 1}段的内容。" for i in range(5)]


def test_wrapped_lines_are_merged_into_one_paragraph(tmp_path: Path) -> None:
    """★ 场景 6（开发刚修的点）：同一段的折行**应该**合并成一块。

    否则下游按句切分时会切到半句话。
    """
    path = _one_page_pdf(
        tmp_path / "wrapped.pdf",
        [
            (72, 100, "这是一段很长的正文，排版时被折成了两行，"),
            (72, 100 + LINE_STEP, "合起来才是一句完整的话。"),
        ],
    )

    doc = parse_pdf(path)
    assert _paragraphs(doc) == ["这是一段很长的正文，排版时被折成了两行，合起来才是一句完整的话。"]
    block = doc.blocks[0]
    assert (block.line_start, block.line_end) == (1, 2)
    assert block.page_no == 1


def test_three_wrapped_lines_become_one_block(tmp_path: Path) -> None:
    """折成三行也要合成一块。"""
    path = _one_page_pdf(
        tmp_path / "wrapped3.pdf",
        [
            (72, 100, "第一行的内容，"),
            (72, 100 + LINE_STEP, "第二行的内容，"),
            (72, 100 + 2 * LINE_STEP, "第三行收尾。"),
        ],
    )
    assert _paragraphs(parse_pdf(path)) == ["第一行的内容，第二行的内容，第三行收尾。"]


def test_paragraph_ending_with_sentence_punctuation_is_not_swallowed(tmp_path: Path) -> None:
    """★ 场景 6：上一段以句末标点收尾 → 两段必须断开，哪怕它们挤在同一个 text block 里。

    最小复现：两段相邻、行距等于单倍行距，上一段以句末标点收尾。

        y=100  "第一段结束了。"
        y=112  "第二段紧接着开始。"      # 单倍行距

    当前行为（= 期望）：两块 —— 同一 text block 内的多行会按段落规则切分
    （`_group_lines` 先把块拆到行，再用与跨块**同一套** `_is_continuation` 判据
    合并），所以"上一行以句末标点收尾 → 下一行是新段落"这条判据在块内同样生效，
    **段末标点即断段**。唯一的豁免是"上一行排满了本块右边界"（`_is_full_line`）——
    那种断行是折行而不是段落边界；本用例两行都远未排满，不受豁免影响。
    """
    path = _one_page_pdf(
        tmp_path / "sentence_end_close.pdf",
        [
            (72, 100, "第一段结束了。"),
            (72, 100 + LINE_STEP, "第二段紧接着开始。"),
        ],
    )

    assert _paragraphs(parse_pdf(path)) == ["第一段结束了。", "第二段紧接着开始。"]


# ---------------------------------------------------------------------------
# 场景 7：中英混排的拼接
# ---------------------------------------------------------------------------


def test_latin_wrap_is_joined_with_exactly_one_space(tmp_path: Path) -> None:
    """★ 场景 7：拉丁文折行 → 只插**一个**空格，不多不少。

    开发接口文档里写的回归点是"合起来才是一句完整的话"，如果行尾/行首的空格
    处理不当就会出现 `line  and`（两个空格）或 `linelayer`（没有空格）。
    """
    path = _one_page_pdf(
        tmp_path / "latin_wrap.pdf",
        [
            (72, 100, "The routing table keeps the next hop for each network"),
            (72, 100 + LINE_STEP, "and the forwarding engine uses it to switch packets."),
        ],
    )
    text = _paragraphs(parse_pdf(path))[0]
    assert "network and the forwarding" in text
    assert "  " not in text, f"出现了连续两个空格：{text!r}"
    assert "networkand" not in text


def test_cjk_wrap_is_joined_without_any_space(tmp_path: Path) -> None:
    """★ 场景 7：中文折行 → 不插空格（`网络` + `层` = `网络层`）。"""
    path = _one_page_pdf(
        tmp_path / "cjk_wrap.pdf",
        [
            (72, 100, "路由器根据路由表决定下一跳，"),
            (72, 100 + LINE_STEP, "下一跳指向到达目的网络的方向。"),
        ],
    )
    text = _paragraphs(parse_pdf(path))[0]
    assert text == "路由器根据路由表决定下一跳，下一跳指向到达目的网络的方向。"
    assert " " not in text


def test_cjk_adjacent_to_latin_gets_no_space(tmp_path: Path) -> None:
    """★ 场景 7：中→英 / 英→中 相邻时**不插**空格。

    行为：`这是一段中文正文被折行` + `layer 继续写。` → `…折行layer 继续写。`
    `_join_lines` 的规则是"只在前后都是非 CJK 时插空格"，所以 CJK↔拉丁 一律
    不插。这条规则**单一、可预测、且不会凭空多插空格**（不会出现 `折行 layer`
    这种多出来的空格），只是拼接结果把中文和英文单词贴在一起了。
    中英之间是否要补一个空格属于产品口径，本批保持现规则。
    """
    path = _one_page_pdf(
        tmp_path / "cjk_latin.pdf",
        [
            (72, 100, "这是一段中文正文被折行"),
            (72, 100 + LINE_STEP, "layer 继续写。"),
        ],
    )
    text = _paragraphs(parse_pdf(path))[0]
    assert text == "这是一段中文正文被折行layer 继续写。"
    assert "折行 layer" not in text, "不该在中英之间凭空多插空格"


def test_latin_adjacent_to_cjk_gets_no_space(tmp_path: Path) -> None:
    """英→中 同样不插空格（与上一条对称）。"""
    path = _one_page_pdf(
        tmp_path / "latin_cjk.pdf",
        [
            (72, 100, "This is a wrapped english line"),
            (72, 100 + LINE_STEP, "网络层继续写。"),
        ],
    )
    assert _paragraphs(parse_pdf(path))[0] == "This is a wrapped english line网络层继续写。"


def test_digits_adjacent_to_cjk_are_kept_verbatim(tmp_path: Path) -> None:
    """★ 场景 7：CJK 与数字相邻（`端口号是` + `65535`）不能凭空多插空格。"""
    path = _one_page_pdf(
        tmp_path / "digits.pdf",
        [
            (72, 100, "端口号是"),
            (72, 100 + LINE_STEP, "65535 这个值。"),
        ],
    )
    text = _paragraphs(parse_pdf(path))[0]
    assert text == "端口号是65535 这个值。"
    assert "是 65535" not in text


def test_mixed_cjk_latin_paragraph_keeps_all_original_characters(tmp_path: Path) -> None:
    """中英混排整段：除了**行间断点**，字符一个都不能丢或改。"""
    path = _one_page_pdf(
        tmp_path / "mixed.pdf",
        [
            (72, 100, "TCP 的可靠传输依赖 sequence number 与 ACK，"),
            (72, 100 + LINE_STEP, "重传由 timeout 触发。"),
        ],
    )
    text = _paragraphs(parse_pdf(path))[0]
    for fragment in ("TCP", "sequence number", "ACK", "timeout", "可靠传输"):
        assert fragment in text
    assert "  " not in text


# ---------------------------------------------------------------------------
# 场景 8：空页 / 空白页
# ---------------------------------------------------------------------------


def test_blank_page_produces_no_block(tmp_path: Path) -> None:
    """★ 场景 8：完全空白的页不能产出块。"""
    doc = pymupdf.open()
    p1 = doc.new_page()
    _w(p1, 72, 100, "第一页有内容。")
    doc.new_page()  # 完全空白
    p3 = doc.new_page()
    _w(p3, 72, 100, "第三页有内容。")
    path = tmp_path / "blank_middle.pdf"
    doc.save(str(path))
    doc.close()

    parsed = parse_pdf(path)
    assert [b.content_md for b in parsed.blocks] == ["第一页有内容。", "第三页有内容。"]
    assert [b.page_no for b in parsed.blocks] == [1, 3]
    assert parsed.page_count == 3


def test_whitespace_only_page_produces_no_block(tmp_path: Path) -> None:
    """★ 场景 8：只有空格/空白的页也不能产出空白块 —— 空白块会在下游造出空卡片。"""
    doc = pymupdf.open()
    p1 = doc.new_page()
    _w(p1, 72, 100, "首页内容。")
    p2 = doc.new_page()
    p2.insert_text((72, 100), "     ", fontsize=10)
    path = tmp_path / "whitespace.pdf"
    doc.save(str(path))
    doc.close()

    parsed = parse_pdf(path)
    assert [b.content_md for b in parsed.blocks] == ["首页内容。"]
    assert all(b.content_md.strip() for b in parsed.blocks)


def test_no_block_has_blank_content_md(tmp_path: Path) -> None:
    """★ 场景 8（通用契约）：任何情况下都不许出现 `content_md.strip() == ""` 的块。"""
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page()
        page.insert_text((72, 100), "   ", fontsize=10)
    page = doc.new_page()
    _w(page, 72, 100, "有内容的页。")
    path = tmp_path / "mostly_blank.pdf"
    doc.save(str(path))
    doc.close()

    parsed = parse_pdf(path)
    assert parsed.blocks
    assert all(b.content_md.strip() for b in parsed.blocks)
    assert all("\n\n\n" not in b.content_md for b in parsed.blocks)


def test_blank_page_at_the_end_does_not_add_a_trailing_block(tmp_path: Path) -> None:
    """尾页空白不能在 Markdown 尾巴上留一个空块。"""
    doc = pymupdf.open()
    p1 = doc.new_page()
    _w(p1, 72, 100, "唯一一页的内容。")
    doc.new_page()
    path = tmp_path / "blank_tail.pdf"
    doc.save(str(path))
    doc.close()

    parsed = parse_pdf(path)
    assert len(parsed.blocks) == 1
    assert parsed.page_count == 2


# ---------------------------------------------------------------------------
# 场景 11：DOCX
# ---------------------------------------------------------------------------


def make_docx_with_structure(path: Path) -> Path:
    """标题（样式 + 多级编号法）+ 正文 + 表格 + 正文：覆盖文档顺序。"""
    document = docx.Document()
    document.add_heading("第1章 网络层", level=1)
    document.add_paragraph("路由是把分组从源送到目的的过程。")
    document.add_heading("1.1 路由基础", level=2)

    numbered = document.add_paragraph("2.3.4 用编号法判出来的三级标题")
    numbered.style = document.styles["Normal"]

    single = document.add_paragraph("3. 单级编号只是列表项，不是标题")
    single.style = document.styles["Normal"]

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "协议"
    table.cell(0, 1).text = "作用"
    table.cell(1, 0).text = "IP"
    table.cell(1, 1).text = "寻址"

    document.add_paragraph("表格后面还有一段正文。")
    document.save(str(path))
    return path


def test_docx_heading_levels_from_style_and_number(tmp_path: Path) -> None:
    """★ 场景 11：样式标题 1 级 / 2 级；多级编号（2.3.4）判为 3 级。"""
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    headings = {b.content_md: b.heading_level for b in doc.blocks if b.block_type == "heading"}
    assert headings["第1章 网络层"] == 1
    assert headings["1.1 路由基础"] == 2
    assert headings["2.3.4 用编号法判出来的三级标题"] == 3


def test_docx_single_level_number_is_not_a_heading(tmp_path: Path) -> None:
    """★ 场景 11：`3. 单级编号…` 在 Word 里几乎总是列表项，不能变成"章"。"""
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    by_text = {b.content_md: b for b in doc.blocks}
    assert by_text["3. 单级编号只是列表项，不是标题"].block_type == "paragraph"
    assert by_text["3. 单级编号只是列表项，不是标题"].heading_level is None


def test_docx_document_order_is_preserved(tmp_path: Path) -> None:
    """★ 场景 11：段落与表格必须保持文档顺序（表格夹在正文中间是常态）。"""
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    assert [b.block_type for b in doc.blocks] == [
        "heading",
        "paragraph",
        "heading",
        "heading",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert doc.blocks[-1].content_md == "表格后面还有一段正文。"


def test_docx_page_no_is_none_for_every_block(tmp_path: Path) -> None:
    """★ 场景 11：DOCX 没有页的概念 —— 一律 None，不许按段落数"估"一个页码。"""
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    assert doc.page_count is None
    assert all(b.page_no is None for b in doc.blocks)


def test_docx_table_renders_as_markdown_table(tmp_path: Path) -> None:
    """★ 场景 11：表格 → Markdown 表格，列数一致、单元格内容不丢。"""
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    table = next(b for b in doc.blocks if b.block_type == "table")
    lines = table.content_md.splitlines()
    assert lines[0] == "| 协议 | 作用 |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| IP | 寻址 |"
    assert all(line.count("|") == 3 for line in lines)


def test_docx_nested_table_content_is_not_lost(tmp_path: Path) -> None:
    """★ 场景 11：嵌套表格的内容**不许静默丢失**。

    最小复现：外层表 2×2，其右下角单元格里再嵌一个 2×2 表。

    当前行为（= 期望）：嵌套表的内容（"内层1"…"内层4"）会被**递归展平**进外层
    单元格的文本里 —— `_cell_content` 按 XML 顺序取单元格的直接段落与嵌套表，
    `_flatten_table` 把内层压成一行（格内空格分隔、行间「/」分隔），四段文字因此
    都在某个块的 `content_md` 里。同时登记一条**中文存疑说明**
    （`ambiguous_structure` / severity=low），讲明"Markdown 表达不了嵌套结构，
    内容已展平写进外层单元格，表格层次需人工核对" —— 丢的只是表格层次，
    内容不丢，而且不吭声是不允许的。
    """
    document = docx.Document()
    document.add_paragraph("外层说明段落。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "外层A"
    table.cell(0, 1).text = "外层B"
    table.cell(1, 0).text = "外层C"
    inner = table.cell(1, 1).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "内层1"
    inner.cell(0, 1).text = "内层2"
    inner.cell(1, 0).text = "内层3"
    inner.cell(1, 1).text = "内层4"
    path = tmp_path / "nested.docx"
    document.save(str(path))

    doc = parse_docx(path)
    joined = "\n".join(b.content_md for b in doc.blocks)
    for cell_text in ("内层1", "内层2", "内层3", "内层4"):
        assert cell_text in joined, f"嵌套表格内容 {cell_text!r} 静默丢失"


def test_docx_without_headings_reports_ambiguous_structure(tmp_path: Path) -> None:
    """没有标题样式 → 骨架推不出来，必须老实给一条中文存疑说明。"""
    document = docx.Document()
    document.add_paragraph("这份材料只有正文。")
    document.add_paragraph("正文第二段。")
    path = tmp_path / "plain.docx"
    document.save(str(path))

    doc = parse_docx(path)
    assert doc.blocks
    assert [n.kind for n in doc.uncertain_notes] == ["ambiguous_structure"]
    assert _HAN_RE.search(doc.uncertain_notes[0].message)


def test_docx_document_with_headings_has_no_structure_note(tmp_path: Path) -> None:
    doc = parse_docx(make_docx_with_structure(tmp_path / "struct.docx"))
    assert doc.uncertain_notes == []


# ---------------------------------------------------------------------------
# 共用块契约（PDF / DOCX 一起扫一遍）
# ---------------------------------------------------------------------------


def test_shared_block_contract_holds_for_pdf_and_docx(tmp_path: Path) -> None:
    """块字段的硬约束：类型枚举、heading_level 规则、非空内容。

    这条是"共用契约"的兜底：无论走哪个解析器，落库前都必须满足
    `material_blocks` 的 CHECK 与 `ParsedBlock.__post_init__` 的约定。
    """
    pdf_path = _one_page_pdf(
        tmp_path / "contract.pdf",
        [
            (72, 72, "第1章 网络层", 18),
            (72, 110, "1.1 路由基础", 13),
            (72, 140, "正文一段。"),
            (72, 140 + PARAGRAPH_GAP, "正文二段。"),
        ],
    )
    docs = [
        parse_pdf(pdf_path),
        parse_docx(make_docx_with_structure(tmp_path / "contract.docx")),
    ]

    for doc in docs:
        assert doc.blocks
        assert [seq for seq, _ in doc.numbered()] == list(range(len(doc.blocks)))
        for block in doc.blocks:
            assert block.block_type in BLOCK_TYPES
            assert block.content_md.strip(), f"出现空白块：{block!r}"
            if block.block_type == "heading":
                assert block.heading_level is not None
                assert 1 <= block.heading_level <= 6
            else:
                assert block.heading_level is None, "非 heading 块带上了 heading_level"
            assert block.page_no is None or block.page_no >= 1


@pytest.mark.parametrize("size", [7.0, 10.0, 14.0, 22.0, 48.0])
def test_heading_detection_never_produces_out_of_range_level(tmp_path: Path, size: float) -> None:
    """字号极端时也不能产出越界的 heading_level（DB 有 1–6 的 CHECK）。"""
    path = _one_page_pdf(
        tmp_path / f"size{int(size)}.pdf",
        [
            (72, 100, "第1章 大字号标题", size),
            (72, 200, "正文内容放在这里。", 10),
        ],
    )
    for block in parse_pdf(path).blocks:
        if block.block_type == "heading":
            assert 1 <= block.heading_level <= 6
        else:
            assert block.heading_level is None
