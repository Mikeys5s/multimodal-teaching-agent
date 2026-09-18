"""DOCX 解析的测试（归属：P1）。SPEC §5.1。

守的是这几条契约：
  · 标题样式 → `block_type="heading"` 且 `heading_level` 正确（1–6）
  · 表格 → `block_type="table"`（内容渲染成 Markdown 表格）
  · `page_no` **必须为 None** —— DOCX 没有页的概念，不许按段落数"估"一个页码
  · `seq` 是**整份材料内**的全局序号，从 0 开始连续（不是节内序号）
  · 段落与表格保持文档顺序（表格夹在段落中间时不能张冠李戴）
  · A2-7 同输入可复现

夹具全部用 python-docx **现场生成**，不依赖任何外部素材文件 —— 素材文件进不了
版本库，靠它们写的测试在队友机器上必然红。
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ApiError, ErrorCode
from app.models._common import BLOCK_TYPES
from app.models.ids import block_id, hash_bytes, material_id
from app.models.material import Material, MaterialBlock
from app.parse import parse_docx, parse_material
from app.parse.persist import persist_blocks

_HAN_RE = re.compile(r"[\u4e00-\u9fff]")

#: 与 test_ids.py 同口径：material_id 只取哈希前 8 位。
MAT_ID = material_id(hash_bytes(b"docx fixture"))


# ---------------------------------------------------------------------------
# 夹具生成
# ---------------------------------------------------------------------------


def make_docx(path: Path) -> Path:
    """标题 + 段落 + 表格 + 段落：覆盖"表格夹在正文中间"这个常态。"""
    document = docx.Document()

    document.add_heading("第1章 网络层", level=1)
    document.add_paragraph("路由是把分组从源送到目的的过程。")
    document.add_heading("1.1 路由基础", level=2)
    document.add_paragraph("路由器根据路由表决定下一跳。")

    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "协议"
    table.cell(0, 1).text = "作用"
    table.cell(1, 0).text = "IP"
    table.cell(1, 1).text = "寻址"
    table.cell(2, 0).text = "TCP"
    table.cell(2, 1).text = "可靠传输"

    document.add_paragraph("表格后面还有一段正文。")
    document.save(str(path))
    return path


def make_plain_docx(path: Path) -> Path:
    """一个标题样式都没有的文档 —— 骨架推不出来，必须老实说。"""
    document = docx.Document()
    document.add_paragraph("这份材料只有正文，没有任何标题样式。")
    document.add_paragraph("正文第二段。")
    document.save(str(path))
    return path


@pytest.fixture
def fixture_docx(tmp_path: Path) -> Path:
    return make_docx(tmp_path / "讲义.docx")


# ---------------------------------------------------------------------------
# 块类型 / 层级
# ---------------------------------------------------------------------------


def test_docx_source_type_and_text_extract(fixture_docx: Path) -> None:
    """DOCX 走文本抽取，且**没有**页数可言（不许编造 page_count）。"""
    doc = parse_docx(fixture_docx)
    assert doc.source_type == "docx"
    assert doc.parse_method == "text_extract"
    assert doc.page_count is None
    assert doc.blocks


def test_headings_keep_text_and_get_their_levels(fixture_docx: Path) -> None:
    """★ 标题：`Heading 1` / `Heading 2` 要落成 1 级与 2 级，文本保持原貌。"""
    doc = parse_docx(fixture_docx)
    by_text = {b.content_md: b for b in doc.blocks if b.block_type == "heading"}
    assert by_text["第1章 网络层"].heading_level == 1
    assert by_text["1.1 路由基础"].heading_level == 2


def test_heading_level_is_none_for_non_heading_blocks(fixture_docx: Path) -> None:
    """非 heading 块带上 level 会让下游凭空多出一节，必须在解析层就为 None。"""
    doc = parse_docx(fixture_docx)
    headings = [b for b in doc.blocks if b.block_type == "heading"]
    assert headings
    for block in doc.blocks:
        if block.block_type == "heading":
            assert block.heading_level is not None
            assert 1 <= block.heading_level <= 6
        else:
            assert block.heading_level is None


def test_every_block_type_is_in_the_enum(fixture_docx: Path) -> None:
    doc = parse_docx(fixture_docx)
    assert {b.block_type for b in doc.blocks} <= set(BLOCK_TYPES)
    assert {b.block_type for b in doc.blocks} == {"heading", "paragraph", "table"}


def test_plain_paragraph_is_not_a_heading(fixture_docx: Path) -> None:
    doc = parse_docx(fixture_docx)
    by_text = {b.content_md: b for b in doc.blocks}
    assert by_text["路由是把分组从源送到目的的过程。"].block_type == "paragraph"


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------


def test_table_becomes_a_markdown_table(fixture_docx: Path) -> None:
    doc = parse_docx(fixture_docx)
    tables = [b for b in doc.blocks if b.block_type == "table"]
    assert len(tables) == 1
    md = tables[0].content_md
    assert "| 协议 | 作用 |" in md
    assert "| --- | --- |" in md
    assert "| IP | 寻址 |" in md
    assert "| TCP | 可靠传输 |" in md


def test_document_order_survives_tables_between_paragraphs(fixture_docx: Path) -> None:
    """`document.paragraphs` 与 `document.tables` 是两个列表，拼起来会丢顺序。"""
    doc = parse_docx(fixture_docx)
    types = [b.block_type for b in doc.blocks]
    assert types == ["heading", "paragraph", "heading", "paragraph", "table", "paragraph"]
    assert doc.blocks[-1].content_md == "表格后面还有一段正文。"


def test_table_block_has_no_page_no(fixture_docx: Path) -> None:
    doc = parse_docx(fixture_docx)
    table = next(b for b in doc.blocks if b.block_type == "table")
    assert table.page_no is None


# ---------------------------------------------------------------------------
# 页码 / 序号
# ---------------------------------------------------------------------------


def test_page_no_is_none_for_every_block(fixture_docx: Path) -> None:
    """★ DOCX 无页概念：页码一律 None，不许"估"。"""
    doc = parse_docx(fixture_docx)
    assert all(b.page_no is None for b in doc.blocks)


def test_seq_is_global_continuous_from_zero(fixture_docx: Path) -> None:
    """★ seq 是整份材料的全局序号：0,1,2…，跨章节也不重置。"""
    doc = parse_docx(fixture_docx)
    assert [seq for seq, _ in doc.numbered()] == list(range(len(doc.blocks)))


def test_block_ids_are_deterministic_from_seq(fixture_docx: Path) -> None:
    doc = parse_docx(fixture_docx)
    expected = [block_id(MAT_ID, seq) for seq in range(len(doc.blocks))]
    assert expected[0].startswith("blk_")
    assert len(set(expected)) == len(expected)


# ---------------------------------------------------------------------------
# 不确定处
# ---------------------------------------------------------------------------


def test_document_without_any_heading_says_so(tmp_path: Path) -> None:
    """没有标题样式 → 骨架推不出来，必须给一条能展示的中文说明。"""
    doc = parse_docx(make_plain_docx(tmp_path / "plain.docx"))
    assert doc.blocks
    assert len(doc.uncertain_notes) == 1
    note = doc.uncertain_notes[0]
    assert note.kind == "ambiguous_structure"
    assert _HAN_RE.search(note.message)
    assert "标题" in note.message


def test_heading_document_has_no_structure_note(fixture_docx: Path) -> None:
    assert parse_docx(fixture_docx).uncertain_notes == []


# ---------------------------------------------------------------------------
# 失败隔离（A1-7）：损坏 / 加密都要是中文
# ---------------------------------------------------------------------------


def test_corrupted_docx_raises_chinese_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is definitely not a docx")
    with pytest.raises(ApiError) as excinfo:
        parse_docx(path)

    err = excinfo.value
    assert err.code == ErrorCode.INVALID_PARAM
    assert _HAN_RE.search(err.message), "报错必须是中文，不能是英文堆栈"
    assert "损坏" in err.message


def test_encrypted_or_legacy_docx_raises_chinese_error(tmp_path: Path) -> None:
    """OLE2 头 = 加密的 .docx 或旧版 .doc，处置动作相同，合并成一条提示。"""
    path = tmp_path / "encrypted.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    with pytest.raises(ApiError) as excinfo:
        parse_docx(path)

    err = excinfo.value
    assert _HAN_RE.search(err.message)
    assert "加密" in err.message


def test_parse_material_dispatches_docx(fixture_docx: Path) -> None:
    assert parse_material(fixture_docx).source_type == "docx"


# ---------------------------------------------------------------------------
# 可复现（A2-7）与落库
# ---------------------------------------------------------------------------


def test_same_docx_parses_identically_twice(fixture_docx: Path) -> None:
    assert parse_docx(fixture_docx) == parse_docx(fixture_docx)


def test_persist_blocks_writes_rows_with_null_page_no(fixture_docx: Path, session: Session) -> None:
    """落库：页码留空、音频字段不写（D-08 保留不启用）、seq 与块锚点同源。"""
    session.add(
        Material(
            id=MAT_ID,
            filename="讲义.docx",
            file_hash=hash_bytes(b"docx fixture"),
            stored_path="uploads/讲义.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size_bytes=fixture_docx.stat().st_size,
            source_type="docx",
            parse_method="text_extract",
        )
    )
    session.flush()

    doc = parse_docx(fixture_docx)
    ids = persist_blocks(session, MAT_ID, doc)

    rows = (
        session.execute(
            select(MaterialBlock)
            .where(MaterialBlock.material_id == MAT_ID)
            .order_by(MaterialBlock.seq)
        )
        .scalars()
        .all()
    )
    assert [row.seq for row in rows] == list(range(len(doc.blocks)))
    assert [row.id for row in rows] == ids
    assert ids == [block_id(MAT_ID, seq) for seq in range(len(doc.blocks))]
    for row in rows:
        assert row.page_no is None
        assert row.ts_start_ms is None and row.ts_end_ms is None
        assert row.ocr_confidence is None
