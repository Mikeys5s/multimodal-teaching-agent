"""QA 独立验证 · 场景 11 边缘：DOCX 表格保真。

本文件守住两条**已按产品口径修正**的表格渲染行为（这里钉住正确行为，
防止回归）：

  1. **横向合并只写一次**：`python-docx` 的 `row.cells` 对横向合并的左单元格
     会重复返回同一个 `Cell` 对象，曾经的渲染把"跨三列"写进 Markdown 3 次 ——
     "一个合并格"在 Markdown 里被伪装成"三个内容相同的格"。现在由
     `_row_cells_with_gaps` 只写首格、续格留空，列数仍然一致。
  2. **纯空表格不产块**：1×1 且单元格为空的表曾经渲染成 `|  |` + 分隔行，
     它的 `content_md.strip()` 仍含 `|`，因此躲过空白块过滤、变成一个纯噪声块，
     会进 Markdown、进库、被当成"节内容"投喂给抽取模型。现在整表无内容时
     `_table_to_md` 返回空串，调用方不产块。

其余用例（列数一致、竖线转义、单元格换行压平）是不变的保真要求。
"""

from __future__ import annotations

import re
from pathlib import Path

import docx

from app.parse import parse_docx

#: 未转义的竖线 = 列分隔符（`\|` 属于单元格内容）
_UNESCAPED_PIPE_RE = re.compile(r"(?<!\\)\|")


def make_merged_cell_docx(path: Path) -> Path:
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "表头A"
    table.cell(0, 1).text = "表头B"
    table.cell(0, 2).text = "表头C"
    merged = table.cell(1, 0)
    merged.merge(table.cell(1, 2))  # 横向合并三列
    merged.text = "跨三列"
    document.save(str(path))
    return path


def test_horizontally_merged_cell_text_is_written_once(tmp_path: Path) -> None:
    """横向合并单元格的内容只写**一次**，其余续格留空（列数不变）。

    Markdown 表格语法没有 colspan，"跨三列"这个语义无法完整表达。实现的取舍是
    "首格写内容 + 续格留空"，而**不是**把同一句重复填进每个续格：重复会让下游
    无法区分"这是一个合并格"和"这里真的写了三遍同样的话"，也会让原文的体积
    与统计口径（字符数、覆盖率）虚高。
    """
    doc = parse_docx(make_merged_cell_docx(tmp_path / "merged.docx"))
    table = next(b for b in doc.blocks if b.block_type == "table")

    assert table.content_md == (
        "| 表头A | 表头B | 表头C |\n| --- | --- | --- |\n| 跨三列 |  |  |"
    )
    assert table.content_md.count("跨三列") == 1, (
        "横向合并的内容被复制了；续格必须留空，不能重复填写"
    )


def test_merged_cell_rows_keep_consistent_column_count(tmp_path: Path) -> None:
    """列数必须一致，否则 Markdown 表格在多数渲染器里会直接塌掉。"""
    doc = parse_docx(make_merged_cell_docx(tmp_path / "merged.docx"))
    table = next(b for b in doc.blocks if b.block_type == "table")
    for line in table.content_md.splitlines():
        assert line.count("|") == 4, f"列数不齐：{line!r}"


def test_empty_table_produces_no_block(tmp_path: Path) -> None:
    """空表格**不产出任何块** —— `|  |` 这种"半空"表格块是纯噪声。

    曾经的渲染产出 `|  |\\n| --- |`，它在 `strip()` 之后仍含 `|`，因此躲过了
    `ParsedDocument` 的空白块过滤，进 Markdown、进库、并被当成"节内容"投喂给
    抽取模型。现在整表无内容时 `_table_to_md` 返回空串，调用方直接不产块。
    """
    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = ""
    path = tmp_path / "empty_table.docx"
    document.save(str(path))

    doc = parse_docx(path)
    assert doc.blocks == [], (
        f"空表格产出了块：{[(b.block_type, repr(b.content_md)) for b in doc.blocks]}"
    )


def test_table_cells_with_pipe_characters_are_escaped(tmp_path: Path) -> None:
    """单元格里的 `|` 必须转义，否则表格会多出一列。"""
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表达式"
    table.cell(0, 1).text = "含义"
    table.cell(1, 0).text = "a | b"
    table.cell(1, 1).text = "按位或"
    path = tmp_path / "pipes.docx"
    document.save(str(path))

    doc = parse_docx(path)
    table_block = next(b for b in doc.blocks if b.block_type == "table")
    assert r"a \| b" in table_block.content_md
    for line in table_block.content_md.splitlines():
        # 只数**未转义**的竖线：`\|` 是单元格内容，不是列分隔符
        assert len(_UNESCAPED_PIPE_RE.findall(line)) == 3, f"列数不齐（转义失败）：{line!r}"


def test_multi_line_cell_text_is_flattened_to_one_line(tmp_path: Path) -> None:
    """单元格内换行会破坏 Markdown 表格结构 —— 必须压成一行。"""
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "第一行"
    table.cell(0, 0).add_paragraph("第二行")
    table.cell(0, 1).text = "另一列"
    path = tmp_path / "multiline.docx"
    document.save(str(path))

    doc = parse_docx(path)
    table_block = next(b for b in doc.blocks if b.block_type == "table")
    assert len(table_block.content_md.splitlines()) == 2, "单元格换行泄漏成了新行"
    assert "第一行 第二行" in table_block.content_md
